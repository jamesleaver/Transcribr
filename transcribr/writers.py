"""Writing a transcript out: .docx, .pdf, .txt, and the extra formats."""

from pathlib import Path
import os

from .speakers import UNATTRIBUTED_LABEL, _export_speaker_colour
from .platformutil import _write_json_atomic
from .paragraphs import _disclaimer_text, format_timestamp, render


def write_paragraphs_to_file(paragraphs, out_path, *, show_timestamp=True,
                              title=None, output_format="txt",
                              speakers=None, verified_by=None,
                              audio_path=None, hide_ts=None,
                              speaker_slots=None, colour_speakers=False,
                              template="plain", indistinct=None):
    """Single entry point for writing transcript output.

    Centralises the txt-vs-docx-vs-pdf switch so the worker (direct-write
    path) and the GUI (review-screen path) can both call the same function.
    `verified_by` switches the appended disclaimer to a named
    certification; `audio_path` is embedded in .docx metadata so
    playback works when the transcript is re-opened later.

    With `colour_speakers` and a `speaker_slots` list (parallel to
    paragraphs, holding the review pane's "1".."9" slot for each), .docx
    and .pdf paragraphs are shaded in that speaker's review-pane colour.
    Plain text has nowhere to put a colour, so it ignores both.

    `template` ("plain", "court", "erisp") restyles a .docx; PDF and
    plain text keep the plain layout. `indistinct` ({paragraph index:
    [words each [indistinct] replaced]}) is kept in a .docx so the
    review pane can show them again when it is reopened.

    Raises ImportError with a friendly message if .docx or .pdf is
    requested but the needed package isn't installed.
    """
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    if output_format == "docx":
        try:
            _write_docx(paragraphs, out_path,
                        show_timestamp=show_timestamp, title=title,
                        speakers=speakers, verified_by=verified_by,
                        audio_path=audio_path, hide_ts=hide_ts,
                        speaker_slots=speaker_slots,
                        colour_speakers=colour_speakers,
                        template=template, indistinct=indistinct)
        except ImportError:
            raise ImportError(
                "Cannot write .docx: the 'python-docx' package is not "
                "installed. Install it with:\n  pip install python-docx\n"
                "Or pick the .txt format instead.")
    elif output_format == "pdf":
        try:
            _write_pdf(paragraphs, out_path,
                       show_timestamp=show_timestamp, title=title,
                       speakers=speakers, verified_by=verified_by,
                       hide_ts=hide_ts, speaker_slots=speaker_slots,
                       colour_speakers=colour_speakers)
        except ImportError:
            raise ImportError(
                "Cannot write .pdf: the 'reportlab' package is not "
                "installed. Install it with:\n  pip install reportlab\n"
                "Or pick the .txt or .docx format instead.")
    else:
        text = render(paragraphs, show_timestamp=show_timestamp,
                      title=title, speakers=speakers, hide_ts=hide_ts)
        out_path.write_text(
            text + "\n" + _disclaimer_text(verified_by) + "\n",
            encoding="utf-8",
        )


# ---- The .docx style set ---------------------------------------------
#
# Everything the writer emits is carried by a named paragraph style, and
# every one of those styles is defined here rather than inherited from
# whatever template Word happens to open the file with. Two reasons.
#
# The document looks the same everywhere. A style the file does not
# define is resolved by the reader's Normal, so a transcript written
# here and opened on a machine with a different default font came out
# in that font.
#
# And it gives the reader a handle. Shading a paragraph directly leaves
# someone who wants a different colour - or no colour - selecting
# paragraphs and clearing formatting by hand, which takes the indents
# with it. Because the colour lives on "Transcribr Speaker 1" and
# "Transcribr Text 1" instead, changing it is one edit in Word's Styles
# pane and every paragraph of that speaker follows. Right-clicking the
# style also offers "Select All Instances", so one speaker's evidence
# can be pulled out of a long transcript in a click.

DOCX_FONT = "Courier New"
DOCX_FONT_PT = 10

# Style names. The first two are also what _parse_docx_transcript looks
# for, along with the pre-0.9.15 names it still recognises.
DOCX_STYLE_TITLE = "Transcribr Title"
DOCX_STYLE_TEXT = "Transcribr Text"
DOCX_STYLE_SPEAKER = "Transcribr Speaker"
DOCX_STYLE_DISCLAIMER = "Transcribr Disclaimer"
DOCX_STYLE_FOOTER = "Transcribr Footer"


def _docx_fix_font(style, name, *, size_pt=None, bold=None, italic=None):
    """Pin a style's font for every script, not just Latin.

    style.font.name sets w:rFonts/@w:ascii only. Word falls back to the
    theme font for anything outside that range, so a transcript with a
    macron or a Greek character would render part of a line in a
    different face."""
    from docx.oxml.ns import qn
    from docx.shared import Pt

    style.font.name = name
    if size_pt is not None:
        style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def _docx_shade_style(style, fill):
    """Give a paragraph style a background colour. python-docx has no
    API for shading, so the element goes into the style's paragraph
    properties: <w:shd w:val="clear" w:color="auto" w:fill="RRGGBB"/>.
    Word fills the whole text column, which is what makes a run of one
    speaker's paragraphs read as a single band."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    ppr = style.element.get_or_add_pPr()
    for old in ppr.findall(qn("w:shd")):
        ppr.remove(old)
    if not fill:
        return
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill.lstrip("#").upper())
    ppr.append(shd)


def _docx_prune_unused_styles(doc):
    """Drop the styles python-docx's blank template carries but this
    document never uses.

    The template ships about 170 of them - Light Shading Accent 3, Book
    Title, nine heading levels. None are used here, and they bury the
    handful that are: someone looking for "Transcribr Speaker 1" to
    change a colour should find it near the top of the Styles pane, not
    three screens down a list of Word's defaults.

    Conservative on purpose. A style is kept if the document applies it,
    if another kept style is built on it or links to it, or if it is a
    document default. Best-effort: a transcript must never fail to save
    over a tidier Styles pane."""
    from docx.oxml.ns import qn

    try:
        styles_el = doc.styles.element
        used = set()
        # Everything the body, headers and footers actually apply.
        parts = [doc.element.body]
        for section in doc.sections:
            for part in (section.header, section.footer,
                         section.first_page_header, section.first_page_footer,
                         section.even_page_header, section.even_page_footer):
                try:
                    parts.append(part._element)
                except Exception:
                    pass
        for part in parts:
            for tag in ("w:pStyle", "w:rStyle", "w:tblStyle"):
                for el in part.iter(qn(tag)):
                    val = el.get(qn("w:val"))
                    if val:
                        used.add(val)

        by_id = {}
        for st in styles_el.findall(qn("w:style")):
            sid = st.get(qn("w:styleId"))
            if sid:
                by_id[sid] = st
                if st.get(qn("w:default")) == "1":
                    used.add(sid)

        # Follow basedOn / next / link so a kept style's ancestry
        # survives with it.
        pending = list(used)
        while pending:
            st = by_id.get(pending.pop())
            if st is None:
                continue
            for tag in ("w:basedOn", "w:next", "w:link"):
                ref = st.find(qn(tag))
                if ref is not None:
                    val = ref.get(qn("w:val"))
                    if val and val not in used:
                        used.add(val)
                        pending.append(val)

        for sid, st in by_id.items():
            if sid not in used:
                styles_el.remove(st)
    except Exception:
        pass


# Document templates for .docx (doc_template setting). "plain" is the
# layout above; the other two restyle it after it is built, so the
# styles the reader recognises on reopening are untouched - only their
# fonts and spacing change, and the page gains numbered lines and a
# running header.
DOC_TEMPLATES = ("plain", "court", "erisp")

_TEMPLATE_SPEC = {
    # A court transcript: a book face, double spaced, every line
    # numbered from 1 on each page so a passage can be cited as
    # "page 4, lines 12-18", and the matter at the head of each page.
    "court": {"font": "Times New Roman", "size": 12, "spacing": 2.0,
              "left_cm": 3.5, "banner": None},
    # A record of interview: typewriter face, 1.5 spaced, numbered lines
    # and the familiar heading. Pairs with the Q&A layout.
    "erisp": {"font": DOCX_FONT, "size": 11, "spacing": 1.5,
              "left_cm": 3.0, "banner": "RECORD OF INTERVIEW"},
}


def _docx_apply_template(doc, template, *, title=None, audio_path=None):
    """Restyle a built transcript for `template` ("court" or "erisp");
    "plain" or anything unknown leaves it as it is."""
    spec = _TEMPLATE_SPEC.get(template)
    if spec is None:
        return
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    # Every paragraph style the writer made, in the template's face.
    # Sizes are set on Normal and the body text; headings keep their
    # own relative weight.
    for style in doc.styles:
        if style.type != WD_STYLE_TYPE.PARAGRAPH:
            continue
        _docx_fix_font(style, spec["font"])
    normal = doc.styles["Normal"]
    _docx_fix_font(normal, spec["font"], size_pt=spec["size"])
    normal.paragraph_format.line_spacing = spec["spacing"]
    for name in (DOCX_STYLE_TEXT, DOCX_STYLE_SPEAKER, DOCX_STYLE_DISCLAIMER,
                 DOCX_STYLE_FOOTER):
        try:
            doc.styles[name].font.size = Pt(spec["size"])
        except KeyError:
            pass

    for section in doc.sections:
        section.left_margin = Cm(spec["left_cm"])
        # Numbered lines, restarting on each page; Word draws them in
        # the margin and leaves headers and footers unnumbered.
        sect = section._sectPr
        for old in sect.findall(qn("w:lnNumType")):
            sect.remove(old)
        ln = OxmlElement("w:lnNumType")
        ln.set(qn("w:countBy"), "1")
        ln.set(qn("w:restart"), "newPage")
        ln.set(qn("w:distance"), str(int(Cm(0.6).twips)))
        # lnNumType must come before pgNumType/cols/docGrid etc.;
        # inserting after pgMar keeps the schema order Word expects.
        pgmar = sect.find(qn("w:pgMar"))
        if pgmar is not None:
            pgmar.addnext(ln)
        else:
            sect.append(ln)

        header = section.header
        header.is_linked_to_previous = False
        p = header.paragraphs[0]
        p.text = ""
        try:
            p.style = doc.styles[DOCX_STYLE_FOOTER]
        except KeyError:
            pass
        if spec["banner"]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(spec["banner"])
            r.bold = True
            if title:
                p2 = header.add_paragraph(title, style=p.style)
                p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            # Heading on the left, the recording it came from on the
            # right.
            width = section.page_width - section.left_margin - section.right_margin
            p.paragraph_format.tab_stops.add_tab_stop(
                width, WD_TAB_ALIGNMENT.RIGHT)
            p.add_run(title or "")
            source = os.path.basename(str(audio_path)) if audio_path else ""
            # With no heading set, the heading is the file name already.
            if source and source != (title or "").strip():
                p.add_run("\t" + source)



def _docx_build_styles(doc, *, show_timestamp, slots, colour_speakers):
    """Define the document's whole style set and return
    {"text": name, "speaker": name, "text_for": {slot: name},
     "speaker_for": {slot: name}, ...}.

    `slots` is the set of speaker slots the document actually uses, so a
    two-speaker transcript gets two pairs of styles rather than all
    nine cluttering the Styles pane."""
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Cm, Pt

    # Normal is the root every other style inherits from. Setting it is
    # what stops the reader's own Normal deciding how this document
    # looks.
    normal = doc.styles["Normal"]
    _docx_fix_font(normal, DOCX_FONT, size_pt=DOCX_FONT_PT)
    npf = normal.paragraph_format
    npf.space_before = Pt(0)
    npf.space_after = Pt(0)
    npf.line_spacing = 1.0
    npf.widow_control = True

    def add(name, *, base=None, quick=False):
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = base or normal
        style.quick_style = quick
        return style

    title = add(DOCX_STYLE_TITLE, quick=True)
    _docx_fix_font(title, DOCX_FONT, size_pt=12, bold=True)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True

    text = add(DOCX_STYLE_TEXT, quick=True)
    tpf = text.paragraph_format
    if show_timestamp:
        # Hanging indent: the timestamp sits in the left margin column
        # and the body wraps to where its first line began.
        tpf.left_indent = Cm(3.0)
        tpf.first_line_indent = -Cm(3.0)
    tpf.space_after = Pt(6)

    speaker = add(DOCX_STYLE_SPEAKER, quick=True)
    _docx_fix_font(speaker, DOCX_FONT, size_pt=DOCX_FONT_PT, bold=True)
    spf = speaker.paragraph_format
    spf.space_before = Pt(8)
    spf.space_after = Pt(0)
    # Keep the label on the same page as the paragraph it introduces;
    # without it Word can strand a name at the foot of a page.
    spf.keep_with_next = True
    speaker.next_paragraph_style = text

    disclaimer = add(DOCX_STYLE_DISCLAIMER)
    _docx_fix_font(disclaimer, DOCX_FONT, size_pt=DOCX_FONT_PT, italic=True)
    disclaimer.paragraph_format.space_before = Pt(12)

    footer = add(DOCX_STYLE_FOOTER)
    _docx_fix_font(footer, DOCX_FONT, size_pt=DOCX_FONT_PT)

    # One pair per speaker actually present. They exist whether or not
    # colouring is on: an uncoloured transcript can then be colour-coded
    # after the fact by giving two styles a fill, without touching a
    # single paragraph.
    text_for, speaker_for = {}, {}
    for slot in sorted(slots, key=lambda s: (len(str(s)), str(s))):
        fill = _export_speaker_colour(slot) if colour_speakers else None
        st = add(f"{DOCX_STYLE_TEXT} {slot}", base=text, quick=True)
        _docx_shade_style(st, fill)
        sl = add(f"{DOCX_STYLE_SPEAKER} {slot}", base=speaker, quick=True)
        _docx_shade_style(sl, fill)
        sl.next_paragraph_style = st
        text_for[slot] = st.name
        speaker_for[slot] = sl.name

    return {
        "title": title.name, "text": text.name, "speaker": speaker.name,
        "disclaimer": disclaimer.name, "footer": footer.name,
        "text_for": text_for, "speaker_for": speaker_for,
    }


# ----- Custom document properties ---------------------------------------
#
# Word keeps custom properties through an ordinary edit, which makes them
# the place for what a layout cannot show: whose each Q&A row was
# (qa.QA_PROPERTY), and the words each [indistinct] replaced.

INDISTINCT_PROPERTY = "TranscribrIndistinct"

_CUSTOM_CT = ("application/vnd.openxmlformats-officedocument."
              "custom-properties+xml")
_CUSTOM_RT = ("http://schemas.openxmlformats.org/officeDocument/2006/"
              "relationships/custom-properties")
_CUSTOM_NS = ("http://schemas.openxmlformats.org/officeDocument/2006/"
              "custom-properties")
_VT_NS = ("http://schemas.openxmlformats.org/officeDocument/2006/"
          "docPropsVTypes")


def _set_custom_properties(doc, props):
    """Write string custom document properties {name: value} (python-docx
    has no API for these). All in one part, which a document may only
    have one of. Best-effort: a transcript never fails to save over it."""
    if not props:
        return
    try:
        from docx.opc.packuri import PackURI
        from docx.opc.part import Part
        from xml.sax.saxutils import escape, quoteattr

        items = "".join(
            '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" '
            f'pid="{pid}" name={quoteattr(name)}>'
            f'<vt:lpwstr>{escape(value)}</vt:lpwstr></property>'
            for pid, (name, value) in enumerate(props.items(), start=2))
        xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
               f'<Properties xmlns="{_CUSTOM_NS}" xmlns:vt="{_VT_NS}">'
               f'{items}</Properties>').encode("utf-8")
        package = doc.part.package
        part = Part(PackURI("/docProps/custom.xml"), _CUSTOM_CT, xml,
                    package)
        package.relate_to(part, _CUSTOM_RT)
    except Exception:
        pass


def read_custom_property(doc, name):
    """A string custom document property, or None."""
    from xml.etree import ElementTree as ET
    for rel in doc.part.package.rels.values():
        if rel.reltype != _CUSTOM_RT:
            continue
        root = ET.fromstring(rel.target_part.blob)
        for prop in root.findall(f"{{{_CUSTOM_NS}}}property"):
            if prop.get("name") == name:
                val = prop.find(f"{{{_VT_NS}}}lpwstr")
                return val.text if val is not None else None
    return None


def _docx_embed_meta(doc, out_path, audio_path, verified_by):
    """Embed Transcribr's small metadata record (source audio, verifier)
    in the document's comments property. Best-effort, never raises."""
    try:
        import json as _json
        base = {}
        if verified_by:
            base["verified_by"] = str(verified_by)
        audio = str(audio_path) if audio_path else None
        rel = None
        if audio:
            try:
                rel = os.path.relpath(
                    audio, os.path.dirname(str(out_path)) or ".")
            except ValueError:      # e.g. different drive on Windows
                pass
        candidates = []
        for extra in ({"audio": audio, "audio_rel": rel},
                      {"audio_rel": rel},
                      {"audio": audio},
                      {}):
            meta = dict(base)
            meta.update({k: v for k, v in extra.items() if v})
            if meta and meta not in candidates:
                candidates.append(meta)
        for meta in candidates:
            encoded = _json.dumps({"transcribr": meta})
            if len(encoded) <= 255:
                doc.core_properties.comments = encoded
                break
    except Exception:
        pass


def _write_docx(paragraphs, out_path, *, show_timestamp=True, title=None,
                speakers=None, verified_by=None, audio_path=None,
                hide_ts=None, speaker_slots=None, colour_speakers=False,
                template="plain", indistinct=None):
    """Write paragraphs to a .docx file with a page-numbered footer and an
    italic disclaimer at the end.

    If `title` is given, it appears in bold at the top of the document
    before the transcript paragraphs.

    If `speakers` is given (a list parallel to paragraphs), each paragraph
    can be prefixed by a bold speaker label on its own line. To match
    legal-transcript convention, the label is only emitted when it differs
    from the previous paragraph's speaker (so a string of paragraphs from
    one speaker reads as a single labelled block).

    When `show_timestamp` is True (default), each paragraph is rendered as
    timestamp + tab + body, with a hanging indent so the timestamp sits in
    the left margin column and the body text wraps further right.

    When `show_timestamp` is False, paragraphs are rendered as flat prose
    with no indent.

    With `colour_speakers`, each paragraph (and the label above it) is
    shaded in that speaker's review-pane colour, so a printed transcript
    reads the way the review pane did."""
    from docx import Document
    from docx.shared import Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    from docx.enum.section import WD_ORIENT

    doc = Document()

    # Force A4 portrait. python-docx defaults to US Letter (8.5x11"), so set
    # the page dimensions explicitly. A4 is 21.0 x 29.7 cm.
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        # Slightly tighter margins to leave more room for the indented body.
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # Every style the document uses, defined here rather than left to
    # the reader's template. Only the speakers actually present get a
    # pair of styles.
    used_slots = {s for s in (speaker_slots or []) if s} if speakers else set()
    st = _docx_build_styles(doc, show_timestamp=show_timestamp,
                            slots=used_slots,
                            colour_speakers=colour_speakers)

    if title:
        doc.add_paragraph(title, style=st["title"])

    last_speaker = None
    for i, para in enumerate(paragraphs):
        start = para[0][0]
        # A paragraph is one paragraph: its internal line breaks are
        # invisible in the review pane and mean nothing in the output,
        # but .txt uses blank lines to separate paragraphs, so leaving
        # them in made one format disagree with the others about how
        # many paragraphs the document has. Collapse here, in all
        # three writers, so the saved file says the same thing however
        # it is written.
        body = " ".join(" ".join(seg[2].split()) for seg in para)
        speaker = speakers[i] if speakers and i < len(speakers) else None
        slot = (speaker_slots[i]
                if speaker_slots and i < len(speaker_slots) else None)

        # Emit a speaker label whenever it changes. For named speakers we
        # emit their name; for an unattributed paragraph that follows a
        # named one, emit an explicit "[Unattributed]" marker so re-parsing
        # the file doesn't carry the previous speaker over.
        if speaker and speaker != last_speaker:
            doc.add_paragraph(
                speaker, style=st["speaker_for"].get(slot, st["speaker"]))
        elif speaker is None and last_speaker is not None:
            doc.add_paragraph(UNATTRIBUTED_LABEL, style=st["speaker"])
        last_speaker = speaker

        p = doc.add_paragraph(style=st["text_for"].get(slot, st["text"]))
        hidden = bool(hide_ts and i < len(hide_ts) and hide_ts[i])
        if show_timestamp and not hidden:
            p.add_run(format_timestamp(start) + "\t" + body)
        else:
            p.add_run(body)

    # Italic disclaimer at the end.
    doc.add_paragraph(_disclaimer_text(verified_by), style=st["disclaimer"])

    # Embed Transcribr's own metadata in the document: the source-audio
    # location (so playback works when the transcript is re-opened) and
    # who verified it (so the review pane's Verify field comes back
    # filled in). Document properties are hard-capped at 255 characters,
    # so the record sheds detail until it fits - the relative audio path
    # is kept in preference to the absolute one because it also survives
    # a whole case folder being moved. Metadata is best-effort: a
    # transcript must never fail to save over it.
    _docx_embed_meta(doc, out_path, audio_path, verified_by)

    # Footer: "Page X of Y" right-aligned, in Courier New so it matches
    # the body font. python-docx doesn't expose Word field codes
    # directly so we drop in the raw XML.
    def _add_field(paragraph, instr):
        run = paragraph.add_run()
        f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
        it.text = instr
        f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
        run._r.append(f1); run._r.append(it); run._r.append(f2)

    def _styled_run(paragraph, text):
        return paragraph.add_run(text)

    fp = doc.sections[0].footer.paragraphs[0]
    fp.style = doc.styles[st["footer"]]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _styled_run(fp, "Page ")
    _add_field(fp, " PAGE ")
    _styled_run(fp, " of ")
    _add_field(fp, " NUMPAGES ")

    _docx_apply_template(doc, template, title=title, audio_path=audio_path)
    if indistinct:
        import json as _json
        _set_custom_properties(
            doc, {INDISTINCT_PROPERTY: _json.dumps(indistinct)})
    _docx_prune_unused_styles(doc)
    doc.save(str(out_path))


def _write_pdf(paragraphs, out_path, *, show_timestamp=True, title=None,
               speakers=None, verified_by=None, hide_ts=None,
               speaker_slots=None, colour_speakers=False):
    """Write paragraphs to an A4 PDF, visually matching the .docx output:
    Courier body with a hanging timestamp indent, bold speaker labels that
    stay with their paragraph, an italic disclaimer, and a right-aligned
    'Page X of Y' footer.

    With `colour_speakers`, each paragraph is shaded in that speaker's
    review-pane colour, matching the .docx."""
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfgen import canvas as _pdfcanvas
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape

    page_w, _page_h = A4

    class _NumberedCanvas(_pdfcanvas.Canvas):
        """Two-pass canvas so the footer can say 'Page X of Y'."""

        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._saved_states = []

        def showPage(self):
            self._saved_states.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._saved_states)
            for state in self._saved_states:
                self.__dict__.update(state)
                self.setFont("Courier", 10)
                self.drawRightString(
                    page_w - 2.5 * cm, 1.5 * cm,
                    f"Page {self._pageNumber} of {total}")
                _pdfcanvas.Canvas.showPage(self)
            _pdfcanvas.Canvas.save(self)

    # The hanging-indent column, in characters. Courier 10pt is 6pt per
    # character, so 14 characters = 84pt (~2.96cm) - the wrapped lines
    # indent to exactly where the first line's body starts, because the
    # timestamp itself is padded to the same column with non-breaking
    # spaces. (The .docx reaches the same alignment with a tab stop;
    # reportlab has no tabs inside paragraphs, hence the padding.)
    _INDENT_CHARS = 14
    _indent_pts = _INDENT_CHARS * 6

    body_style = ParagraphStyle(
        "Transcript",
        fontName="Courier", fontSize=10, leading=12,
        spaceAfter=6,
    )
    if show_timestamp:
        # Hanging indent: timestamp in the left column, wrapped body
        # lines aligned under the body column (mirrors the docx style).
        body_style.leftIndent = _indent_pts
        body_style.firstLineIndent = -_indent_pts
    label_style = ParagraphStyle(
        "SpeakerLabel",
        fontName="Courier-Bold", fontSize=10, leading=12,
        spaceBefore=8, spaceAfter=0, keepWithNext=1,
    )
    title_style = ParagraphStyle(
        "TranscriptTitle",
        fontName="Courier-Bold", fontSize=12, leading=14,
        spaceAfter=12,
    )
    disclaimer_style = ParagraphStyle(
        "Disclaimer",
        fontName="Courier-Oblique", fontSize=10, leading=12,
        spaceBefore=12,
    )

    # A shaded variant of a style, made once per (style, colour) pair.
    # backColor fills the paragraph's own box, so the padding is what
    # stops the text sitting hard against the edge of the block; the
    # left padding is widened to cover the hanging timestamp column.
    _shaded_cache = {}

    def _shaded(base, fill):
        if not fill:
            return base
        key = (base.name, fill)
        if key not in _shaded_cache:
            style = ParagraphStyle(
                f"{base.name}-{fill.lstrip('#')}", parent=base,
                backColor=HexColor(fill), borderPadding=(2, 3, 2, 3),
            )
            _shaded_cache[key] = style
        return _shaded_cache[key]

    story = []
    if title:
        story.append(Paragraph(escape(title), title_style))

    last_speaker = None
    for i, para in enumerate(paragraphs):
        start = para[0][0]
        # A paragraph is one paragraph: its internal line breaks are
        # invisible in the review pane and mean nothing in the output,
        # but .txt uses blank lines to separate paragraphs, so leaving
        # them in made one format disagree with the others about how
        # many paragraphs the document has. Collapse here, in all
        # three writers, so the saved file says the same thing however
        # it is written.
        body = " ".join(" ".join(seg[2].split()) for seg in para)
        speaker = speakers[i] if speakers and i < len(speakers) else None
        fill = None
        if colour_speakers and speaker_slots and i < len(speaker_slots):
            fill = _export_speaker_colour(speaker_slots[i])
        if speaker and speaker != last_speaker:
            story.append(Paragraph(escape(speaker),
                                   _shaded(label_style, fill)))
        elif speaker is None and last_speaker is not None:
            story.append(Paragraph(escape(UNATTRIBUTED_LABEL), label_style))
        last_speaker = speaker

        hidden = bool(hide_ts and i < len(hide_ts) and hide_ts[i])
        if show_timestamp and not hidden:
            # Pad the timestamp to the indent column with non-breaking
            # spaces so the first line's body starts exactly where the
            # wrapped lines are indented to (the .docx reaches the same
            # alignment with a tab stop).
            ts = format_timestamp(start)
            pad = " " * max(1, _INDENT_CHARS - len(ts))
            text = escape(ts) + pad + escape(body)
        else:
            text = escape(body)
        story.append(Paragraph(text, _shaded(body_style, fill)))

    story.append(Spacer(1, 6))
    story.append(Paragraph(
        escape(_disclaimer_text(verified_by)),
        disclaimer_style,
    ))

    doc = SimpleDocTemplate(
        str(out_path), pagesize=A4,
        leftMargin=2.5 * cm, rightMargin=2.5 * cm,
        topMargin=2.5 * cm, bottomMargin=2.5 * cm,
        title=title or Path(out_path).stem,
    )
    doc.build(story, canvasmaker=_NumberedCanvas)


def _reviewed_result(paragraphs, result):
    """A result-shaped dict whose segments carry the REVIEWED text, so
    the extra outputs (SRT/VTT/TSV) reflect the user's edits, splits
    and merges rather than the raw engine pass. The JSON extra keeps
    the raw engine result - it is the technical record."""
    segments = [
        {"start": float(seg[0]), "end": float(seg[1]),
         "text": str(seg[2])}
        for para in paragraphs
        for seg in para
        if str(seg[2]).strip()
    ]
    out = dict(result or {})
    out["segments"] = segments
    return out


def _write_extra_formats(result, txt_out_path, formats, q,
                         raw_result=None):
    """Write SRT/VTT/JSON/TSV next to the paragraph .txt.

    Engine-agnostic: reads from result["segments"] which every engine
    runner normalises to a list of {"start", "end", "text"} dicts.
    `raw_result`, when given (the review-save path), is written for the
    JSON extra instead of `result` - the JSON is the raw engine
    record, while SRT/VTT/TSV follow the reviewed text."""
    segments = (result or {}).get("segments", []) if result else []
    if not segments:
        q.put(("log", "Note: no segments available for extra outputs.\n"))
        return

    output_dir = txt_out_path.parent
    stem = txt_out_path.stem
    for suffix in (".transcript", ".paragraphs"):
        if stem.endswith(suffix):
            stem = stem[: -len(suffix)]
            break

    for fmt in formats:
        out = output_dir / f"{stem}.{fmt}"
        try:
            if fmt == "json":
                _write_json_result(out, raw_result or result)
            elif fmt == "srt":
                _write_srt(out, segments)
            elif fmt == "vtt":
                _write_vtt(out, segments)
            elif fmt == "tsv":
                _write_tsv(out, segments)
            else:
                q.put(("log",
                       f"WARNING: unknown extra format '{fmt}', skipped.\n"))
                continue
            q.put(("log", f"Wrote: {out.name}\n"))
        except Exception as e:
            q.put(("log",
                   f"WARNING: failed to write .{fmt}: "
                   f"{type(e).__name__}: {e}\n"))


def _seg_field(seg, key):
    """Pull start/end/text from either a dict or an object segment."""
    if isinstance(seg, dict):
        return seg.get(key)
    return getattr(seg, key, None)


def _format_srt_time(seconds):
    seconds = max(0.0, float(seconds))
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = seconds - h * 3600 - m * 60
    return f"{h:02d}:{m:02d}:{s:06.3f}".replace(".", ",")


def _format_vtt_time(seconds):
    seconds = max(0.0, float(seconds))
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = seconds - h * 3600 - m * 60
    return f"{h:02d}:{m:02d}:{s:06.3f}"


def _write_srt(path, segments):
    with open(path, "w", encoding="utf-8") as fh:
        for i, seg in enumerate(segments, start=1):
            start = _seg_field(seg, "start") or 0
            end = _seg_field(seg, "end") or 0
            text = (_seg_field(seg, "text") or "").strip()
            fh.write(f"{i}\n")
            fh.write(f"{_format_srt_time(start)} --> "
                     f"{_format_srt_time(end)}\n")
            fh.write(f"{text}\n\n")


def _write_vtt(path, segments):
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("WEBVTT\n\n")
        for seg in segments:
            start = _seg_field(seg, "start") or 0
            end = _seg_field(seg, "end") or 0
            text = (_seg_field(seg, "text") or "").strip()
            fh.write(f"{_format_vtt_time(start)} --> "
                     f"{_format_vtt_time(end)}\n")
            fh.write(f"{text}\n\n")


def _write_tsv(path, segments):
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("start\tend\ttext\n")
        for seg in segments:
            start = _seg_field(seg, "start") or 0
            end = _seg_field(seg, "end") or 0
            text = (_seg_field(seg, "text") or "").strip().replace("\t", " ")
            fh.write(f"{int(float(start) * 1000)}\t"
                     f"{int(float(end) * 1000)}\t{text}\n")


def _write_json_result(path, result):
    """Dump the full result dict as JSON. Coerces any non-serialisable
    attributes (e.g. faster-whisper named tuples) to plain dicts."""

    def _coerce(obj):
        if isinstance(obj, (str, int, float, bool, type(None))):
            return obj
        if isinstance(obj, dict):
            return {k: _coerce(v) for k, v in obj.items()}
        if isinstance(obj, (list, tuple)):
            return [_coerce(v) for v in obj]
        if hasattr(obj, "_asdict"):
            return _coerce(obj._asdict())
        if hasattr(obj, "__dict__"):
            return _coerce(vars(obj))
        return str(obj)

    # Atomic like every other write: this sits next to the user's
    # transcript, and a half-written technical record is worse than an
    # absent one. No schema envelope - the file is the engine's own
    # result, and other tools read it.
    _write_json_atomic(path, _coerce(result))
