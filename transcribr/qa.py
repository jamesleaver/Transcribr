"""Question-and-answer layout, numbered the way an ERISP transcript is.

Police ERISP transcripts number each question and mark each answer:

    Q1   What is your full name?
    A    John Citizen.
    Q2   And your date of birth?

This module decides which paragraphs are questions and which answers,
and lays the transcript out that way - in the review pane, where Q&A is
a mode the reviewer can switch on and off, and in the saved document.

Who asks is taken from the speaker labels where the transcript has
them: by default the speakers whose paragraphs read most like questions
are the askers and everyone else answers, and the reviewer can say
otherwise speaker by speaker. Paragraphs with no label are placed by
the two-party role model (transcribr/turns.py), from the wording alone.
The document says which of the two it relied on, because a Q or an A
put against the wrong person in a record of interview is a real error.

When more than one person asks questions, the first question after a
change of asker carries the new asker's name - "Q15 (Det Smith)" - so
the reader can follow who took over. "[No audio ...]" markers are not
speech: they are neither questions nor answers and carry no number.

A saved Q&A .docx carries what the layout cannot show (whose each
paragraph was, and who asks) in a custom document property, so the
review pane can reopen it in Q&A mode with the speakers intact.
"""

import json
import re
from pathlib import Path

from .paragraphs import _disclaimer_text, format_timestamp
from . import turns as _turns

# The custom document property a Q&A .docx carries.
QA_PROPERTY = "TranscribrQA"

_NO_AUDIO_RE = re.compile(r"^\s*\[No audio\b", re.IGNORECASE)


def _para_text(para):
    return " ".join(" ".join(seg[2].split()) for seg in para).strip()


def is_no_audio(para):
    """A "[No audio from A to B]" marker paragraph, not speech."""
    return bool(_NO_AUDIO_RE.match(_para_text(para)))


def suggest_askers(paragraphs, slots):
    """The speaker slots whose paragraphs read most like questions.
    Empty when fewer than two speakers are labelled - there is no one to
    tell apart."""
    score, count = {}, {}
    for para, slot in zip(paragraphs, slots):
        if not slot or is_no_audio(para):
            continue
        text = _para_text(para)
        sents = [s for s in text.replace("?", "?\n").replace(". ", ".\n")
                 .splitlines() if s.strip()] or [text]
        score[slot] = score.get(slot, 0.0) + sum(
            _turns.asks(s) for s in sents) / len(sents)
        count[slot] = count.get(slot, 0) + 1
    if len(score) < 2:
        return []
    mean = {s: score[s] / count[s] for s in score}
    top = max(mean.values())
    if top <= 0:
        return []
    return sorted(s for s, v in mean.items() if v >= 0.6 * top
                  and v > min(mean.values()))


def qa_roles(paragraphs, slots=None, askers=None):
    """("Q" | "A" | None per paragraph, askers used, method).

    method is "labels" when every spoken paragraph's role came from its
    speaker label, "wording" when none did, and "mixed" otherwise.
    No-audio markers get None and are left out of the count."""
    slots = list(slots or [None] * len(paragraphs))
    slots += [None] * (len(paragraphs) - len(slots))
    if askers is None:
        askers = suggest_askers(paragraphs, slots)
    askers = set(askers or [])
    spoken = [i for i, p in enumerate(paragraphs) if not is_no_audio(p)]
    spoken_set = set(spoken)
    units = [(paragraphs[i][0][0], paragraphs[i][-1][1],
              _para_text(paragraphs[i]), []) for i in spoken]
    guessed = {}
    if units:
        for i, (r, _) in zip(spoken, _turns.two_party_roles(
                units, unit="paragraph")):
            guessed[i] = "Q" if r == 0 else "A"
    roles, from_labels = [], 0
    for i, slot in enumerate(slots[:len(paragraphs)]):
        if i not in spoken_set:
            roles.append(None)
        elif slot and askers:
            roles.append("Q" if slot in askers else "A")
            from_labels += 1
        else:
            roles.append(guessed.get(i))
    n = len(spoken)
    method = ("labels" if n and from_labels == n else
              "wording" if from_labels == 0 else "mixed")
    return roles, sorted(askers), method


def qa_marks(roles, slots=None):
    """The mark printed against each paragraph: "Q<n>" where a question
    begins, "A" where an answer begins, "" where the same side carries
    on (or for an unplaced paragraph). Questions are numbered from 1; an
    answer carries no number. With `slots`, a different known asker
    starts a new question even straight after the last one."""
    slots = list(slots or [None] * len(roles))
    marks, n, last, last_slot = [], 0, None, None
    for r, slot in zip(roles, slots + [None] * (len(roles) - len(slots))):
        new_asker = (r == "Q" and last == "Q" and slot and last_slot
                     and slot != last_slot)
        if r == "Q" and (last != "Q" or new_asker):
            n += 1
            marks.append(f"Q{n}")
        elif r == "A" and last != "A":
            marks.append("A")
        else:
            marks.append("")
        last = r
        if r == "Q" and slot:
            last_slot = slot
    return marks


def qa_layout(paragraphs, slots=None, askers=None, names=None):
    """Everything the pane and the writers need to show the Q&A layout:

        rows    - per paragraph {"mark": "Q3" | "A" | "", "asker": name
                  or None, "role": "Q" | "A" | None}; "asker" is set on
                  the first question after a change of questioner, when
                  more than one person asks
        askers  - the speaker slots treated as asking
        method  - "labels" | "wording" | "mixed"
        note    - the sentence saying how the roles were decided
    """
    slots = list(slots or [None] * len(paragraphs))
    slots += [None] * (len(paragraphs) - len(slots))
    names = names or {}
    roles, used, method = qa_roles(paragraphs, slots, askers=askers)
    marks = qa_marks(roles, slots)
    asking = {slots[i] for i, r in enumerate(roles)
              if r == "Q" and slots[i]}
    several = len(asking) > 1
    rows, prev_asker = [], None
    for i, (r, mark) in enumerate(zip(roles, marks)):
        asker = None
        if r == "Q" and mark and slots[i]:
            if several and slots[i] != prev_asker:
                asker = names.get(slots[i]) or f"Speaker {slots[i]}"
            prev_asker = slots[i]
        rows.append({"mark": mark, "asker": asker, "role": r})
    asker_names = [names.get(a) or f"Speaker {a}" for a in used]
    return {"rows": rows, "askers": used, "method": method,
            "note": method_note(method, asker_names)}


def method_note(method, asker_names):
    who = ", ".join(asker_names) if asker_names else "the questioner"
    if method == "labels":
        return (f"Q and A follow the speaker labels: questions by {who}.")
    if method == "wording":
        return ("Q and A were suggested automatically from the wording, not "
                "from who was speaking. Check each against the recording "
                "before relying on it.")
    return (f"Q and A follow the speaker labels where there are any "
            f"(questions by {who}); unlabelled paragraphs were placed "
            f"automatically from the wording. Check those against the "
            f"recording before relying on them.")


def _asker_prefix(name):
    return f"({name}) "


def write_qa(paragraphs, out_path, marks, *, fmt="docx", title=None,
             note=None, show_timestamp=False, hide_ts=None,
             verified_by=None, askers=None, meta=None, audio_path=None,
             template="plain", indistinct=None):
    """Write the Q&A layout as .docx, .pdf or .txt.

    `askers` is a per-paragraph list of the name to show on a question
    where the questioner changes (or None). `meta`, for .docx, is the
    record that lets the review pane reopen the file in Q&A mode."""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    rows = []
    for i, para in enumerate(paragraphs):
        hidden = bool(hide_ts and i < len(hide_ts) and hide_ts[i])
        stamp = (format_timestamp(para[0][0]) + "  "
                 if show_timestamp and not hidden else "")
        name = askers[i] if askers and i < len(askers) else None
        rows.append((marks[i] if i < len(marks) else "",
                     stamp + (_asker_prefix(name) if name else "")
                     + _para_text(para)))
    if fmt == "txt":
        lines = [title, ""] if title else []
        if note:
            lines += [note, ""]
        lines += [f"{m:<6}{body}" for m, body in rows]
        lines += ["", _disclaimer_text(verified_by)]
        out_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        return out_path
    if fmt == "pdf":
        _write_qa_pdf(rows, out_path, title=title, note=note,
                      verified_by=verified_by)
        return out_path
    if fmt != "docx":
        raise ValueError(f"Q&A layout is written as docx, pdf or txt, "
                         f"not {fmt}")
    _write_qa_docx(rows, out_path, title=title, note=note,
                   verified_by=verified_by, meta=meta,
                   audio_path=audio_path, template=template,
                   indistinct=indistinct)
    return out_path


def _write_qa_docx(rows, out_path, *, title, note, verified_by, meta=None,
                   audio_path=None, template="plain", indistinct=None):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm, Pt
    from .writers import (INDISTINCT_PROPERTY, _docx_apply_template,
                          _docx_build_styles, _docx_embed_meta,
                          _docx_prune_unused_styles, _set_custom_properties)

    doc = Document()
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
    st = _docx_build_styles(doc, show_timestamp=False, slots=set(),
                            colour_speakers=False)
    body = doc.styles[st["text"]]
    pf = body.paragraph_format
    # The mark sits in a 1.5 cm column; the text wraps beside it.
    pf.left_indent = Cm(1.5)
    pf.first_line_indent = -Cm(1.5)
    pf.tab_stops.add_tab_stop(Cm(1.5))
    pf.space_after = Pt(6)
    if title:
        doc.add_paragraph(title, style=st["title"])
    if note:
        doc.add_paragraph(note, style=st["disclaimer"])
    for mark, text in rows:
        p = doc.add_paragraph(style=st["text"])
        if mark:
            p.add_run(mark).bold = True
        p.add_run("\t" + text)
    doc.add_paragraph(_disclaimer_text(verified_by), style=st["disclaimer"])
    _docx_apply_template(doc, template, title=title, audio_path=audio_path)
    _docx_prune_unused_styles(doc)
    _docx_embed_meta(doc, out_path, audio_path, verified_by)
    props = {}
    if meta is not None:
        props[QA_PROPERTY] = json.dumps(meta)
    if indistinct:
        props[INDISTINCT_PROPERTY] = json.dumps(indistinct)
    _set_custom_properties(doc, props)
    doc.save(str(out_path))


def _write_qa_pdf(rows, out_path, *, title, note, verified_by):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table
    from reportlab.platypus import TableStyle
    from xml.sax.saxutils import escape

    base = ParagraphStyle("qa", fontName="Courier", fontSize=10,
                          leading=13)
    bold = ParagraphStyle("qab", parent=base, fontName="Courier-Bold")
    head = ParagraphStyle("qah", parent=bold, fontSize=11, spaceAfter=12)
    small = ParagraphStyle("qan", parent=base, fontName="Courier-Oblique",
                           fontSize=8, leading=10, spaceAfter=10)
    story = []
    if title:
        story.append(Paragraph(escape(title), head))
    if note:
        story.append(Paragraph(escape(note), small))
    data = [[Paragraph(escape(m), bold), Paragraph(escape(t), base)]
            for m, t in rows]
    if data:
        table = Table(data, colWidths=[1.5 * cm, None])
        table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(table)
    story.append(Paragraph(escape(_disclaimer_text(verified_by)), small))
    SimpleDocTemplate(str(out_path), pagesize=A4, leftMargin=2.5 * cm,
                      rightMargin=2.5 * cm).build(story)


# The custom document properties live in writers.py, shared with the
# ordinary layout.
from .writers import read_custom_property  # noqa: E402,F401


def split_qa_mark(text):
    """(mark, rest) for one saved Q&A row: the mark column is whatever
    precedes the first tab, when it is a mark or empty."""
    if "\t" in text:
        head, _, rest = text.partition("\t")
        if not head.strip() or re.fullmatch(r"Q\d+|A", head.strip()):
            return head.strip(), rest.strip()
    return "", text.strip()


def strip_asker(text, name):
    """Drop the "(Det Smith) " a row carries where the questioner
    changed."""
    if name and text.startswith(_asker_prefix(name)):
        return text[len(_asker_prefix(name)):]
    return text
