"""[indistinct] suggestions: offered from low-confidence words, applied
only when the reviewer accepts them (transcribr/transcript.py)."""

import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import _harness  # noqa: F401

import transcribr as T


def model(paragraphs, words):
    m = T.TranscriptModel(paragraphs, word_conf=words)
    return m


# One paragraph, two segments; "mumble jumble" is a low run, "frob" a
# single very unsure word, "think" merely low on its own.
PARAS = [[(0.0, 3.0, " I went down to the mumble jumble shop."),
          (3.0, 6.0, " Then frob I think we left.")]]
WORDS = [
    (0.0, 0.3, " I", 0.9), (0.3, 0.6, " went", 0.9),
    (0.6, 0.9, " down", 0.9), (0.9, 1.1, " to", 0.9),
    (1.1, 1.3, " the", 0.9), (1.3, 1.8, " mumble", 0.2),
    (1.8, 2.3, " jumble", 0.3), (2.3, 3.0, " shop.", 0.9),
    (3.0, 3.4, " Then", 0.9), (3.4, 3.8, " frob", 0.1),
    (3.8, 4.0, " I", 0.9), (4.0, 4.4, " think", 0.3),
    (4.4, 4.8, " we", 0.9), (4.8, 6.0, " left.", 0.9),
]


class TestSuggestions(unittest.TestCase):

    def test_runs_and_very_unsure_words_are_offered(self):
        m = model(PARAS, WORDS)
        body = m.body(0)
        offered = [body[a:b] for a, b in m.indistinct_suggestions()[0]]
        self.assertEqual(offered, ["mumble jumble", "frob"])

    def test_nothing_offered_without_word_confidence(self):
        m = model(PARAS, None)
        self.assertEqual(m.indistinct_suggestions(), [[]])

    def test_nothing_offered_in_a_paragraph_still_transcribing(self):
        m = model(PARAS, WORDS)
        m.sealed_count = 0
        self.assertEqual(m.indistinct_suggestions(), [[]])


class TestMarking(unittest.TestCase):

    def test_marking_replaces_only_the_passage_and_keeps_segments(self):
        m = model(PARAS, WORDS)
        a, b = m.indistinct_suggestions()[0][0]
        before = m.body(0)
        self.assertTrue(m.mark_indistinct(0, a, b))
        self.assertEqual(m.body(0),
                         before.replace("mumble jumble", "[indistinct]"))
        # Still two segments with their own times.
        self.assertEqual([(s, e) for s, e, _ in m.paragraphs[0]],
                         [(0.0, 3.0), (3.0, 6.0)])
        self.assertEqual(m.text_origin[0], "user")

    def test_the_other_offers_survive_a_marking(self):
        """The words a marker replaced must not throw the alignment of
        the words after it - that would quietly withdraw every other
        offer, and the confidence shading with it."""
        m = model(PARAS, WORDS)
        a, b = m.indistinct_suggestions()[0][0]
        m.mark_indistinct(0, a, b)
        body = m.body(0)
        offered = [body[x:y] for x, y in m.indistinct_suggestions()[0]]
        self.assertEqual(offered, ["frob"])
        low = [body[x:y] for x, y, lv in m.confidence_spans()[0]
               if lv == "low"]
        self.assertEqual(low, ["frob", "think"])

    def test_marking_across_segments_joins_them(self):
        m = model(PARAS, WORDS)
        body = m.body(0)
        a = body.index("shop.")
        b = body.index("Then") + len("Then")
        self.assertTrue(m.mark_indistinct(0, a, b))
        self.assertEqual(m.body(0),
                         "I went down to the mumble jumble [indistinct] "
                         "frob I think we left.")
        self.assertEqual([(s, e) for s, e, _ in m.paragraphs[0]],
                         [(0.0, 6.0)])

    def test_undo_restores_the_words(self):
        m = model(PARAS, WORDS)
        a, b = m.indistinct_suggestions()[0][0]
        before = m.body(0)
        m.mark_indistinct(0, a, b)
        m.undo()
        self.assertEqual(m.body(0), before)

    def test_a_kept_passage_is_not_offered_again(self):
        m = model(PARAS, WORDS)
        a, b = m.indistinct_suggestions()[0][0]
        self.assertTrue(m.keep_passage(0, a, b))
        body = m.body(0)
        self.assertEqual(
            [body[x:y] for x, y in m.indistinct_suggestions()[0]],
            ["frob"])

    def test_words_that_moved_are_refused(self):
        """An offer shown before an edit must not land on whatever now
        sits at its old offsets."""
        m = model(PARAS, WORDS)
        a, b = m.indistinct_suggestions()[0][0]
        self.assertFalse(m.mark_indistinct(0, a, b, expect="other words"))
        self.assertTrue(m.mark_indistinct(0, a, b, expect="mumble jumble"))

    def test_bad_offsets_are_refused(self):
        m = model(PARAS, WORDS)
        self.assertFalse(m.mark_indistinct(0, 5, 500))
        self.assertFalse(m.mark_indistinct(3, 0, 1))


class TestSession(unittest.TestCase):

    def test_payload_offers_and_the_action_applies(self):
        info = {"paragraphs": PARAS, "out_path": "/tmp/x.transcript.txt",
                "show_timestamp": True, "title": None,
                "output_format": "txt", "result": {"segments": []},
                "extra_formats": [], "loaded": False, "audio_path": None,
                "word_conf": WORDS}
        s = T.ReviewSession.from_fresh(info, T.EventBroker())
        para = s.payload()["paragraphs"][0]
        self.assertEqual(len(para["unsure"]), 2)
        a, b = para["unsure"][1]
        full = s.mutate(s.model.rev, "indistinct",
                        {"index": 0, "start": a, "end": b})
        self.assertIn("Then [indistinct] I think",
                      full["paragraphs"][0]["body"])
        with self.assertRaises(T.ApiFail):
            s.mutate(full["rev"], "indistinct",
                     {"index": 0, "start": 0, "end": 999})



class TestRestoring(unittest.TestCase):
    """An accepted marker remembers the words it replaced, shows them,
    and can put them back."""

    def _marked(self):
        m = model(PARAS, WORDS)
        for a, b in reversed(m.indistinct_suggestions()[0]):
            m.mark_indistinct(0, a, b)
        return m

    def test_each_marker_carries_its_words(self):
        m = self._marked()
        self.assertEqual([w for _, _, w, _ in m.marker_spans(0)],
                         ["mumble jumble", "frob"])
        body = m.body(0)
        for a, b, _, _ in m.marker_spans(0):
            self.assertEqual(body[a:b], "[indistinct]")

    def test_restoring_puts_the_words_back_and_keeps_the_rest(self):
        m = self._marked()
        self.assertTrue(m.restore_indistinct(0, 1, expect="frob"))
        self.assertIn("Then frob I think", m.body(0))
        self.assertEqual([w for _, _, w, _ in m.marker_spans(0)],
                         ["mumble jumble"])
        # Still two segments with their own times.
        self.assertEqual([(s, e) for s, e, _ in m.paragraphs[0]],
                         [(0.0, 3.0), (3.0, 6.0)])

    def test_restore_refuses_words_it_was_not_shown(self):
        m = self._marked()
        self.assertFalse(m.restore_indistinct(0, 0, expect="frob"))
        self.assertFalse(m.restore_indistinct(0, 5))

    def test_undo_brings_the_record_back_with_the_text(self):
        m = self._marked()
        m.restore_indistinct(0, 0)
        m.undo()
        self.assertEqual([w for _, _, w, _ in m.marker_spans(0)],
                         ["mumble jumble", "frob"])

    def test_words_follow_a_split_and_a_join(self):
        m = self._marked()
        body = m.body(0)
        new = m.split(0, body.index("Then"))
        self.assertEqual([w for _, _, w, _ in m.marker_spans(0)],
                         ["mumble jumble"])
        self.assertEqual([w for _, _, w, _ in m.marker_spans(new)], ["frob"])
        m.merge_with_previous(new)
        self.assertEqual([w for _, _, w, _ in m.marker_spans(0)],
                         ["mumble jumble", "frob"])

    def test_after_a_hand_edit_markers_are_plain_text(self):
        m = self._marked()
        m.commit_edit(0, m.body(0).replace("frob", "frob,", 1)
                      if "frob" in m.body(0) else m.body(0) + " Yes.")
        self.assertEqual(m.marker_spans(0), [])
        self.assertFalse(m.restore_indistinct(0, 0))
        m.undo()
        self.assertEqual([w for _, _, w, _ in m.marker_spans(0)],
                         ["mumble jumble", "frob"])

    def test_a_saved_word_file_keeps_the_words(self):
        try:
            import docx  # noqa: F401
        except ImportError:
            self.skipTest("needs python-docx")
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            out = str(Path(tmp) / "int.transcript.docx")
            info = {"paragraphs": PARAS, "out_path": out,
                    "show_timestamp": True, "title": "T",
                    "output_format": "docx", "result": {"segments": []},
                    "extra_formats": [], "loaded": False,
                    "audio_path": None, "word_conf": WORDS}
            s = T.ReviewSession.from_fresh(info, T.EventBroker())
            for a, b in reversed(s.model.indistinct_suggestions()[0]):
                s.mutate(s.model.rev, "indistinct",
                         {"index": 0, "start": a, "end": b})
            s.save(s.model.rev, "labels")
            again = T.ReviewSession.from_fresh(
                T.open_transcript_info(out), T.EventBroker())
            markers = again.payload()["paragraphs"][0]["markers"]
            self.assertEqual([w for _, _, w, _ in markers],
                             ["mumble jumble", "frob"])
            full = again.mutate(again.model.rev, "indistinct-restore",
                                {"index": 0, "ordinal": 0,
                                 "text": "mumble jumble"})
            self.assertIn("the mumble jumble shop",
                          full["paragraphs"][0]["body"])

    def test_offers_can_be_switched_off(self):
        from unittest import mock
        info = {"paragraphs": PARAS, "out_path": "/tmp/x.transcript.txt",
                "show_timestamp": True, "title": None,
                "output_format": "txt", "result": {"segments": []},
                "extra_formats": [], "loaded": False, "audio_path": None,
                "word_conf": WORDS}
        s = T.ReviewSession.from_fresh(info, T.EventBroker())
        with mock.patch("transcribr.review.current_settings",
                        return_value={"offer_indistinct": False}):
            self.assertEqual(s.payload()["paragraphs"][0]["unsure"], [])

if __name__ == "__main__":
    unittest.main()
