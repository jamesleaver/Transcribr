"""Automated test suite for Transcribr.

Run from the project root with:

    python3 -m unittest discover -s tests -v

or on macOS double-click tests/run_tests.command (uses the app's venv,
which has the optional dependencies installed).

Uses only the standard library (unittest). Tests that need an optional
package (python-docx, reportlab) or a working Tk display skip themselves
when it isn't available.

All config-directory access (settings.json, recent.json, autosave.json)
is redirected to a temporary directory for the duration of the run, so
the suite never touches the real user configuration.
"""

import os
import re
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import _harness  # noqa: F401  (redirects the config dir; must precede transcribr)

try:
    import transcribr as T
except ImportError as e:  # e.g. no tkinter in this interpreter
    raise unittest.SkipTest(
        f"cannot import transcribr in this interpreter: {e}")



# =====================================================================
# Pure-logic helpers
# =====================================================================

class TestFormatTimestamp(unittest.TestCase):
    def test_minutes_seconds(self):
        self.assertEqual(T.format_timestamp(0), "[00:00]")
        self.assertEqual(T.format_timestamp(75.9), "[01:15]")

    def test_hours(self):
        self.assertEqual(T.format_timestamp(3600), "[1:00:00]")
        self.assertEqual(T.format_timestamp(3725), "[1:02:05]")


class TestParagraphify(unittest.TestCase):
    # Texts chosen to avoid every break trigger except the one under
    # test: no sentence-ending punctuation, not short responses.
    A = "and then we walked along"
    B = "toward the river for a while"

    def test_no_break_when_continuous(self):
        segs = [(0.0, 1.0, self.A), (1.2, 2.0, self.B)]
        self.assertEqual(len(T.paragraphify(segs, 1.5)), 1)

    def test_gap_breaks(self):
        # Capitalised continuation so the mid-sentence guard stays out
        # of the way and the silence rule alone is under test.
        segs = [(0.0, 1.0, self.A), (5.0, 6.0, "Toward the river then")]
        self.assertEqual(len(T.paragraphify(segs, 1.5)), 2)

    def test_sentence_end_with_pause_breaks(self):
        # A full stop plus a beat of silence (>= 40% of the threshold)
        # is a paragraph boundary...
        segs = [(0.0, 1.0, "It was finished."), (1.8, 2.5, self.B)]
        self.assertEqual(len(T.paragraphify(segs, 1.5)), 2)

    def test_sentence_end_without_pause_flows_on(self):
        # ...but a full stop with no real pause is just a sentence: a
        # monologue read at speed stays one paragraph.
        segs = [(0.0, 1.0, "It was finished."), (1.1, 2.0, self.B)]
        self.assertEqual(len(T.paragraphify(segs, 1.5)), 1)

    def test_monologue_of_sentences_stays_together(self):
        segs = [(i * 2.0, i * 2.0 + 1.9, "It kept going on.")
                for i in range(5)]
        self.assertEqual(len(T.paragraphify(segs, 1.5)), 1)

    def test_question_always_breaks(self):
        segs = [(0.0, 1.0, "What happened next?"), (1.05, 2.0, self.B)]
        self.assertEqual(len(T.paragraphify(segs, 1.5)), 2)

    def test_interruption_ending_breaks(self):
        for ending in ("I was going to—", "I was going to...",
                       "I was going to--"):
            segs = [(0.0, 1.0, ending), (1.05, 2.0, self.B)]
            self.assertEqual(len(T.paragraphify(segs, 1.5)), 2, ending)

    def test_leading_acknowledgment_needs_a_pause(self):
        # "Yeah, ..." after a beat of silence marks an answer...
        segs = [(0.0, 1.0, self.A), (1.8, 3.0, "Yeah I thought so too")]
        self.assertEqual(len(T.paragraphify(segs, 1.5)), 2)
        # ...but with no pause it's the same speaker's flow.
        segs = [(0.0, 1.0, self.A), (1.05, 3.0, "Yeah I thought so too")]
        self.assertEqual(len(T.paragraphify(segs, 1.5)), 1)

    def test_short_response_breaks(self):
        segs = [(0.0, 1.0, self.A), (1.1, 1.5, "Yes"), (1.6, 2.5, self.B)]
        paras = T.paragraphify(segs, 5.0)
        self.assertEqual(len(paras), 3)

    def test_sixty_second_cap(self):
        # 13 segments, 10s apart, no other break trigger; the cap should
        # split them into three paragraphs (0-50s, 60-110s, 120s).
        segs = [(i * 10.0, i * 10.0 + 9.0, self.A) for i in range(13)]
        paras = T.paragraphify(segs, 100.0)
        self.assertEqual(len(paras), 3)

    def test_never_breaks_mid_sentence(self):
        # Real-case regression: "...on the 27th of March" | "2026. What
        # is your name?" was split by a hard silence. A break must not
        # land mid-sentence, whatever the pause.
        segs = [(0.0, 2.0, "It happened on the 27th of March"),
                (4.5, 6.0, "2026. What is your name?")]
        self.assertEqual(len(T.paragraphify(segs, 1.5)), 1)
        segs = [(0.0, 2.0, "and then he told me that I was"),
                (4.5, 6.0, "behaving badly and had to leave")]
        self.assertEqual(len(T.paragraphify(segs, 1.5)), 1)

    def test_uppercase_start_after_missing_punctuation_can_break(self):
        # Whisper sometimes drops the final full stop; a capitalised
        # fresh sentence after a long silence is still a boundary.
        segs = [(0.0, 2.0, "and that was the end of it"),
                (4.5, 6.0, "Detective Jones entered the room")]
        self.assertEqual(len(T.paragraphify(segs, 1.5)), 2)

    def test_cap_still_splits_endless_continuation(self):
        # The 60s cap remains the last resort even for run-ons.
        segs = [(i * 10.0, i * 10.0 + 9.0, "and it went on and on")
                for i in range(13)]
        self.assertEqual(len(T.paragraphify(segs, 1.5)), 3)

    def test_wall_to_wall_segments_break_on_punctuation_alone(self):
        # faster-whisper without word timestamps pads segments
        # wall-to-wall (end == next start), so no pause is ever
        # visible; sentence endings must then break unconditionally.
        segs = [(0.0, 4.4, "A pencil writes best."),
                (4.4, 9.4, "The lamp shone brightly."),
                (9.4, 12.2, "Clothes are free to new men."),
                (12.2, 14.6, "The glow deepened after dark.")]
        self.assertEqual(len(T.paragraphify(segs, 1.5)), 4)

    def test_empty(self):
        self.assertEqual(T.paragraphify([], 1.5), [])


class TestRefineSegmentTimes(unittest.TestCase):
    def test_padding_trimmed_to_words(self):
        # Whisper padded the segment out to 0.0-4.0 but speech spans
        # 0.8-3.1; the gap measurements should see the real silence.
        segs = [(0.0, 4.0, "hello there"), (4.0, 8.0, "and welcome")]
        words = [(0.8, 1.5, " hello", None), (1.6, 3.1, " there", None),
                 (4.9, 5.5, " and", None), (5.6, 7.2, " welcome", None)]
        out = T._refine_segment_times(segs, words)
        self.assertEqual(out[0][:2], (0.8, 3.1))
        self.assertEqual(out[1][:2], (4.9, 7.2))
        self.assertEqual(out[0][2], "hello there")

    def test_segment_without_words_passes_through(self):
        segs = [(0.0, 2.0, "no words"), (2.0, 4.0, "has words")]
        words = [(2.2, 3.8, " has words", None)]
        out = T._refine_segment_times(segs, words)
        self.assertEqual(out[0], segs[0])
        self.assertEqual(out[1][:2], (2.2, 3.8))

    def test_only_snaps_inward(self):
        # A word timestamp that strays outside its segment is ignored.
        segs = [(1.0, 2.0, "x")]
        words = [(0.5, 2.5, " x", None)]
        self.assertEqual(T._refine_segment_times(segs, words), segs)

    def test_no_words_is_identity(self):
        segs = [(0.0, 1.0, "a")]
        self.assertEqual(T._refine_segment_times(segs, []), segs)


class TestLabelParagraphs(unittest.TestCase):
    A = "and then we walked along"
    B = "toward the river for a while"

    def _para(self, start, end, text):
        return [(start, end, text)]

    def test_two_speaker_conversation(self):
        paragraphs = [self._para(0.0, 8.0, self.A),
                      self._para(10.0, 18.0, self.B)]
        turns = [(0.0, 8.0, 0), (10.0, 18.0, 1)]
        self.assertEqual(T.label_paragraphs(paragraphs, turns),
                         ["1", "2"])

    def test_first_voice_becomes_speaker_one(self):
        paragraphs = [self._para(0.0, 4.0, self.A),
                      self._para(5.0, 9.0, self.B)]
        # Diarizer ids in reverse order: id 7 speaks first.
        turns = [(0.0, 4.0, 7), (5.0, 9.0, 2)]
        self.assertEqual(T.label_paragraphs(paragraphs, turns),
                         ["1", "2"])

    def test_low_confidence_left_unlabelled(self):
        # The turn covers only the last fifth of the paragraph.
        paragraphs = [self._para(0.0, 10.0, self.A)]
        turns = [(8.0, 10.0, 0)]
        self.assertEqual(T.label_paragraphs(paragraphs, turns), [None])

    def test_majority_speaker_wins_a_mixed_paragraph(self):
        # 70% speaker 0 / 30% speaker 1 -> clear majority, labelled.
        paragraphs = [self._para(0.0, 10.0, self.A)]
        turns = [(0.0, 7.0, 0), (7.0, 10.0, 1)]
        self.assertEqual(T.label_paragraphs(paragraphs, turns), ["1"])

    def test_close_call_between_voices_left_unlabelled(self):
        # 55/45 between two voices is not a clear majority.
        paragraphs = [self._para(0.0, 10.0, self.A)]
        turns = [(0.0, 5.5, 0), (5.5, 10.0, 1)]
        self.assertEqual(T.label_paragraphs(paragraphs, turns), [None])

    def test_padded_paragraph_still_labels_its_one_voice(self):
        # Engines pad spans with silence: only 40% of the paragraph is
        # attributed, but it is all one voice - label it.
        paragraphs = [self._para(0.0, 10.0, self.A)]
        turns = [(3.0, 7.0, 0)]
        self.assertEqual(T.label_paragraphs(paragraphs, turns), ["1"])

    def test_numbering_counts_only_labelled_paragraphs(self):
        # A gated-out paragraph must not consume slot 1: the first
        # LABELLED voice becomes Speaker 1.
        paragraphs = [self._para(0.0, 10.0, self.A),
                      self._para(11.0, 15.0, self.B)]
        turns = [(9.0, 10.0, 0),       # 10% coverage -> gated out
                 (11.0, 15.0, 5)]
        self.assertEqual(T.label_paragraphs(paragraphs, turns),
                         [None, "1"])

    def test_no_turns_no_labels(self):
        paragraphs = [self._para(0.0, 4.0, self.A)]
        self.assertEqual(T.label_paragraphs(paragraphs, []), [None])

    def test_paragraph_boundaries_never_change(self):
        # label_paragraphs takes finished paragraphs and only labels
        # them - callers rely on the list being untouched.
        paragraphs = [self._para(0.0, 4.0, self.A),
                      self._para(5.0, 9.0, self.B)]
        snapshot = [list(p) for p in paragraphs]
        T.label_paragraphs(paragraphs, [(0.0, 9.0, 3)])
        self.assertEqual(paragraphs, snapshot)

    def test_more_than_nine_speakers_keeps_busiest(self):
        paragraphs, turns = [], []
        t = 0.0
        for spk in range(10):
            dur = 10.0 - spk
            paragraphs.append(self._para(t, t + dur, "It stops here."))
            turns.append((t, t + dur, spk))
            t += dur + 3.0
        letters = T.label_paragraphs(paragraphs, turns)
        self.assertEqual(letters[:9], [str(i) for i in range(1, 10)])
        self.assertIsNone(letters[9])


class TestVoiceModelRegistry(unittest.TestCase):
    def test_registry_is_sane(self):
        ids = [m["id"] for m in T.DIARIZE_VOICE_MODELS]
        targets = [m["target"] for m in T.DIARIZE_VOICE_MODELS]
        self.assertEqual(len(ids), len(set(ids)), "duplicate model ids")
        self.assertEqual(len(targets), len(set(targets)),
                         "two models would overwrite each other's file")
        self.assertIn(T.DIARIZE_DEFAULT_VOICE_MODEL, ids)
        self.assertNotIn(T._DIARIZE_SEGMENTATION["target"], targets)
        for m in T.DIARIZE_VOICE_MODELS:
            self.assertTrue(m["url"].startswith("https://"), m["id"])
            self.assertEqual(len(m["sha256"]), 64, m["id"])
            for key in ("label", "note", "size", "target"):
                self.assertTrue(m.get(key), f"{m['id']} missing {key}")

    def test_unknown_id_falls_back_to_default(self):
        spec = T._voice_model_spec("nonexistent-model")
        self.assertEqual(spec["id"], T.DIARIZE_DEFAULT_VOICE_MODEL)

    def test_validate_settings_rejects_unknown_voice_model(self):
        merged = T.validate_settings({"diarize_model": "bogus"})
        self.assertEqual(merged["diarize_model"],
                         T.DIARIZE_DEFAULT_VOICE_MODEL)
        merged = T.validate_settings({"diarize_model": "wespeaker"})
        self.assertEqual(merged["diarize_model"], "wespeaker")
        # CAM++ was withdrawn; a stored choice of it must not survive.
        merged = T.validate_settings({"diarize_model": "campplus"})
        self.assertEqual(merged["diarize_model"],
                         T.DIARIZE_DEFAULT_VOICE_MODEL)


class TestDiarizeModelDownload(unittest.TestCase):
    def _spec(self, url, sha, member=None, target="model.onnx"):
        return {"url": url, "sha256": sha, "archive_member": member,
                "target": target, "label": "test model"}

    def _sha(self, path):
        import hashlib
        return hashlib.sha256(Path(path).read_bytes()).hexdigest()

    def test_plain_download_and_verify(self):
        import queue
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "weights.bin"
            src.write_bytes(b"onnx" * 1000)
            spec = self._spec(src.as_uri(), self._sha(src),
                              target="embed.onnx")
            q = queue.Queue()
            T._download_one_model(spec, q, None)
            got = T._diarize_models_dir() / "embed.onnx"
            self.assertTrue(got.exists())
            self.assertEqual(got.read_bytes(), src.read_bytes())
            got.unlink()

    def test_bad_hash_rejected_and_not_installed(self):
        import queue
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "weights.bin"
            src.write_bytes(b"evil")
            spec = self._spec(src.as_uri(), "0" * 64, target="bad.onnx")
            with self.assertRaises(T.DiarizationUnavailable):
                T._download_one_model(spec, queue.Queue(), None)
            self.assertFalse(
                (T._diarize_models_dir() / "bad.onnx").exists())

    def test_tarball_member_extracted(self):
        import queue
        import tarfile
        with tempfile.TemporaryDirectory() as d:
            payload = Path(d) / "model.onnx"
            payload.write_bytes(b"segmentation-weights")
            tar_path = Path(d) / "pack.tar.bz2"
            with tarfile.open(tar_path, "w:bz2") as tar:
                tar.add(payload, arcname="pack/model.onnx")
            spec = self._spec(tar_path.as_uri(), self._sha(tar_path),
                              member="model.onnx", target="seg.onnx")
            T._download_one_model(spec, queue.Queue(), None)
            got = T._diarize_models_dir() / "seg.onnx"
            self.assertEqual(got.read_bytes(), b"segmentation-weights")
            got.unlink()

    def test_existing_target_skips_download(self):
        import queue
        target = T._diarize_models_dir() / "have.onnx"
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(b"already here")
        spec = self._spec("file:///nonexistent/nowhere.onnx", "0" * 64,
                          target="have.onnx")
        T._download_one_model(spec, queue.Queue(), None)   # must not raise
        self.assertEqual(target.read_bytes(), b"already here")
        target.unlink()


class TestDescribeNonAudioFile(unittest.TestCase):
    def _probe(self, data, suffix=".mp3"):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / f"fake{suffix}"
            p.write_bytes(data)
            return T.describe_non_audio_file(str(p))

    def test_office_document_named_mp3(self):
        reason = self._probe(b"PK\x03\x04" + b"\x00" * 64)
        self.assertIn("Microsoft Office", reason)

    def test_pdf_named_mp3(self):
        reason = self._probe(b"%PDF-1.7\n" + b"\x00" * 64)
        self.assertIn("PDF", reason)

    def test_empty_file(self):
        self.assertIn("empty", self._probe(b""))

    def test_missing_file(self):
        reason = T.describe_non_audio_file("/nonexistent/audio.mp3")
        self.assertIn("could not be read", reason)

    def test_real_audio_gets_no_complaint(self):
        try:
            import av  # noqa: F401
        except ImportError:
            self.skipTest("needs the av (PyAV) package")
        import math
        import struct
        import wave
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "tone.wav"
            with wave.open(str(p), "wb") as w:
                w.setnchannels(1)
                w.setsampwidth(2)
                w.setframerate(16000)
                w.writeframes(b"".join(
                    struct.pack("<h", int(8000 * math.sin(i / 20)))
                    for i in range(16000)))
            self.assertIsNone(T.describe_non_audio_file(str(p)))


def _have_av():
    try:
        import av  # noqa: F401
        return True
    except ImportError:
        return False


@unittest.skipUnless(_have_av(), "needs the av (PyAV) package")
class TestPyAvAudio(unittest.TestCase):
    def _make_wav(self, path, seconds=1.0, rate=44100):
        import math
        import struct
        import wave
        with wave.open(str(path), "wb") as w:
            w.setnchannels(1)
            w.setsampwidth(2)
            w.setframerate(rate)
            n = int(seconds * rate)
            frames = b"".join(
                struct.pack("<h", int(12000 * math.sin(
                    2 * math.pi * 440 * i / rate)))
                for i in range(n))
            w.writeframes(frames)

    def test_decode_audio_16k(self):
        with tempfile.TemporaryDirectory() as d:
            wav = Path(d) / "tone.wav"
            self._make_wav(wav, seconds=1.0)
            audio = T.audioprep.decode_audio_16k(wav)
            self.assertIsNotNone(audio)
            self.assertEqual(str(audio.dtype), "float32")
            self.assertAlmostEqual(
                len(audio) / T.AUDIO_SAMPLE_RATE, 1.0, delta=0.05)
            self.assertLessEqual(float(abs(audio).max()), 1.0)

    def test_get_audio_duration(self):
        with tempfile.TemporaryDirectory() as d:
            wav = Path(d) / "tone.wav"
            self._make_wav(wav, seconds=2.0)
            duration = T.get_audio_duration(str(wav))
            self.assertIsNotNone(duration)
            self.assertAlmostEqual(duration, 2.0, delta=0.1)

    def test_decode_missing_file_returns_none(self):
        self.assertIsNone(T.audioprep.decode_audio_16k("/nonexistent/audio.mp3"))


class TestRevisionPath(unittest.TestCase):
    def test_first_revision_and_rev_stripping(self):
        with tempfile.TemporaryDirectory() as d:
            original = Path(d) / "case.transcript.docx"
            original.write_text("x")
            rev1 = T._next_revision_path(original)
            self.assertEqual(rev1.name, "case.transcript.rev1.docx")
            rev1.write_text("x")
            # Next revision skips the existing rev1.
            rev2 = T._next_revision_path(original)
            self.assertEqual(rev2.name, "case.transcript.rev2.docx")
            # A revision of a revision doesn't nest suffixes.
            rev2b = T._next_revision_path(rev1)
            self.assertEqual(rev2b.name, "case.transcript.rev2.docx")


class TestGuessAudio(unittest.TestCase):
    def test_finds_sibling_media(self):
        with tempfile.TemporaryDirectory() as d:
            d = Path(d)
            (d / "interview.mp3").write_bytes(b"")
            transcript = d / "interview.transcript.docx"
            transcript.write_bytes(b"")
            self.assertEqual(
                T._guess_audio_for_transcript(transcript),
                str(d / "interview.mp3"))

    def test_revision_strips_rev_suffix(self):
        with tempfile.TemporaryDirectory() as d:
            d = Path(d)
            (d / "interview.m4a").write_bytes(b"")
            rev = d / "interview.transcript.rev3.docx"
            rev.write_bytes(b"")
            self.assertEqual(
                T._guess_audio_for_transcript(rev),
                str(d / "interview.m4a"))

    def test_no_media_returns_none(self):
        with tempfile.TemporaryDirectory() as d:
            t = Path(d) / "interview.transcript.txt"
            t.write_text("")
            self.assertIsNone(T._guess_audio_for_transcript(t))


class TestExtractWordConf(unittest.TestCase):
    def test_extracts_words(self):
        res = {"segments": [
            {"start": 0, "end": 2, "text": "hi there",
             "words": [
                 {"word": "hi", "start": 0, "end": 1, "probability": 0.9},
                 {"word": " there", "start": 1, "end": 2,
                  "probability": 0.3},
             ]},
            {"start": 2, "end": 3, "text": "no words"},
        ]}
        wc = T._extract_word_conf(res)
        self.assertEqual(len(wc), 2)
        self.assertEqual(wc[1][2].strip(), "there")
        self.assertAlmostEqual(wc[1][3], 0.3)

    def test_handles_missing_probability(self):
        res = {"segments": [
            {"words": [{"word": "x", "start": 0, "end": 1}]}]}
        wc = T._extract_word_conf(res)
        self.assertEqual(len(wc), 1)
        self.assertIsNone(wc[0][3])

    def test_empty_inputs(self):
        self.assertEqual(T._extract_word_conf(None), [])
        self.assertEqual(T._extract_word_conf({}), [])


# =====================================================================
# Config-dir persistence (settings / recent / autosave)
# =====================================================================

class TestTheme(unittest.TestCase):
    def test_resolve_explicit_settings(self):
        self.assertEqual(T._resolve_theme("light"), "light")
        self.assertEqual(T._resolve_theme("dark"), "dark")

    def test_resolve_auto_returns_valid_palette_key(self):
        self.assertIn(T._resolve_theme("auto"), T._PALETTES)

    def test_palette_follows_active_theme(self):
        old = T.palettes._ACTIVE_THEME
        try:
            T.palettes._ACTIVE_THEME = "dark"
            self.assertEqual(T._palette(), T._PALETTES["dark"])
            T.palettes._ACTIVE_THEME = "light"
            self.assertEqual(T._palette(), T._PALETTES["light"])
        finally:
            T.palettes._ACTIVE_THEME = old

    def test_palettes_have_matching_keys(self):
        light, dark = T._PALETTES["light"], T._PALETTES["dark"]
        self.assertEqual(set(light.keys()), set(dark.keys()))
        self.assertEqual(set(light["speaker_colours"]),
                         set(dark["speaker_colours"]))


class TestSettingsPersistence(unittest.TestCase):
    def test_round_trip(self):
        T._settings_save({"model": "tiny", "gap": 2.5, "diarize": True})
        loaded = T._settings_load()
        self.assertEqual(loaded["model"], "tiny")
        self.assertEqual(loaded["gap"], 2.5)
        self.assertIs(loaded["diarize"], True)

    def test_corrupt_file_returns_empty(self):
        T._settings_file().write_text("{not json", encoding="utf-8")
        self.assertEqual(T._settings_load(), {})


class TestRecentList(unittest.TestCase):
    def setUp(self):
        T._recent_save([])

    def test_add_moves_to_front_and_dedupes(self):
        with tempfile.TemporaryDirectory() as d:
            a = Path(d) / "a.docx"; a.write_text("")
            b = Path(d) / "b.docx"; b.write_text("")
            T._recent_add(a)
            T._recent_add(b)
            T._recent_add(a)  # again: moves to front, no duplicate
            items = T._recent_load()
            self.assertEqual(items[0]["path"], str(a.resolve()))
            self.assertEqual(len(items), 2)

    def test_capped_at_max(self):
        for i in range(T._RECENT_MAX + 5):
            T._recent_add(f"/tmp/file{i}.txt")
        self.assertEqual(len(T._recent_load()), T._RECENT_MAX)


class TestAutosave(unittest.TestCase):
    def test_round_trip_and_clear(self):
        data = {
            "out_path": "/tmp/x.docx",
            "paragraphs": [[[0.0, 1.0, "hello"]]],
            "speakers": ["1"],
            "speaker_names": {"1": "Witness"},
        }
        T._autosave_save(data)
        loaded = T._autosave_load()
        self.assertIsNotNone(loaded)
        self.assertEqual(loaded["speakers"], ["1"])
        T._autosave_clear()
        self.assertIsNone(T._autosave_load())

    def test_rejects_empty_paragraphs(self):
        T._autosave_save({"out_path": "/tmp/x.docx", "paragraphs": []})
        self.assertIsNone(T._autosave_load())
        T._autosave_clear()


class TestPersistedFileSchema(unittest.TestCase):
    """The envelope, the migrations, and the atomic write."""

    def test_written_files_carry_the_envelope(self):
        import json
        T._settings_save({"model": "tiny"})
        doc = json.loads(T._settings_file().read_text(encoding="utf-8"))
        self.assertEqual(doc["kind"], "settings")
        self.assertEqual(doc["schema"], T._SCHEMA["settings"])
        self.assertEqual(doc["data"]["model"], "tiny")

    def test_pre_envelope_settings_still_read(self):
        """0.9.14 wrote the bare dict with no envelope."""
        import json
        T._settings_file().write_text(
            json.dumps({"model": "small.en", "gap": 3.0}), encoding="utf-8")
        loaded = T._settings_load()
        self.assertEqual(loaded["model"], "small.en")
        self.assertEqual(loaded["gap"], 3.0)

    def test_pre_0_9_5_recent_strings_are_migrated(self):
        import json
        T._recent_file().write_text(
            json.dumps(["/tmp/a.docx", "/tmp/b.docx"]), encoding="utf-8")
        items = T._recent_load()
        self.assertEqual([i["path"] for i in items],
                         ["/tmp/a.docx", "/tmp/b.docx"])
        self.assertIs(items[0]["reviewed"], False)
        self.assertIs(items[0]["verified"], False)

    def test_pre_envelope_recent_dicts_are_migrated(self):
        import json
        T._recent_file().write_text(
            json.dumps([{"path": "/tmp/a.docx", "reviewed": True,
                         "verified": False}]), encoding="utf-8")
        items = T._recent_load()
        self.assertEqual(len(items), 1)
        self.assertIs(items[0]["reviewed"], True)

    def test_a_newer_schema_is_read_rather_than_discarded(self):
        """A downgrade must not silently wipe the user's settings."""
        import json
        T._settings_file().write_text(json.dumps({
            "schema": T._SCHEMA["settings"] + 5,
            "kind": "settings",
            "data": {"model": "tiny"},
        }), encoding="utf-8")
        self.assertEqual(T._settings_load()["model"], "tiny")

    def test_corrupt_envelope_falls_back_to_default(self):
        T._recent_file().write_text("{not json", encoding="utf-8")
        self.assertEqual(T._recent_load(), [])

    def test_atomic_write_leaves_no_temp_files(self):
        target = T._config_dir() / "atomic-probe.json"
        T._write_json_atomic(target, {"a": 1})
        self.assertEqual(
            list(T._config_dir().glob("atomic-probe.json.*")), [])
        target.unlink()

    def test_failed_write_leaves_the_previous_file_intact(self):
        """A serialisation blow-up must not truncate what was there."""
        target = T._config_dir() / "keepme.json"
        T._write_json_atomic(target, {"good": True})
        before = target.read_text(encoding="utf-8")
        with self.assertRaises(TypeError):
            T._write_json_atomic(target, {"bad": object()})
        self.assertEqual(target.read_text(encoding="utf-8"), before)
        self.assertEqual(list(T._config_dir().glob("keepme.json.*")), [])
        target.unlink()

    def test_every_kind_round_trips(self):
        for kind in T._SCHEMA:
            path = T._config_dir() / f"probe-{kind}.json"
            T._write_versioned(path, kind, {"v": kind})
            self.assertEqual(T._read_versioned(path, kind), {"v": kind})
            path.unlink()

    def test_every_migration_step_is_reachable(self):
        """A migration keyed above the current schema would never run."""
        for kind, steps in T._MIGRATIONS.items():
            self.assertIn(kind, T._SCHEMA)
            for from_version in steps:
                self.assertLess(from_version, T._SCHEMA[kind])


class TestSpeakerDetectionDefaults(unittest.TestCase):
    """The speaker-detection defaults changed. Saving settings writes
    every value, defaults
    included, so the change has to be carried to existing settings by a
    migration - and only to values nobody chose."""

    def setUp(self):
        self.addCleanup(lambda: T._settings_file().unlink(missing_ok=True))

    def _write(self, data, schema=1):
        import json
        doc = data if schema == 0 else {
            "schema": schema, "kind": "settings", "data": data}
        T._settings_file().write_text(json.dumps(doc), encoding="utf-8")

    def test_fresh_install_defaults(self):
        d = T.validate_settings({})
        self.assertEqual(d["diarize_model"], "titanet")
        self.assertEqual(d["diarize_model"], T.DIARIZE_DEFAULT_VOICE_MODEL)
        self.assertEqual(d["num_speakers"], 2)
        self.assertEqual(d["diarize_threshold"], 0.9)

    def test_default_is_offered_first_and_campplus_is_not(self):
        ids = [m["id"] for m in T.DIARIZE_VOICE_MODELS]
        self.assertEqual(ids[0], T.DIARIZE_DEFAULT_VOICE_MODEL)
        self.assertNotIn("campplus", ids)

    def test_old_defaults_move_to_the_new_ones(self):
        self._write({"diarize_model": "wespeaker", "num_speakers": 0,
                     "diarize_threshold": 0.5, "model": "small.en"})
        s = T.current_settings()
        self.assertEqual(s["diarize_model"], "titanet")
        self.assertEqual(s["num_speakers"], 2)
        self.assertEqual(s["diarize_threshold"], 0.9)
        self.assertEqual(s["model"], "small.en")    # nothing else moves

    def test_deliberate_choices_are_left_alone(self):
        self._write({"num_speakers": 3, "diarize_threshold": 0.7})
        s = T.current_settings()
        self.assertEqual(s["num_speakers"], 3)
        self.assertEqual(s["diarize_threshold"], 0.7)

    def test_withdrawn_campplus_becomes_the_default(self):
        self._write({"diarize_model": "campplus"})
        self.assertEqual(T.current_settings()["diarize_model"], "titanet")

    def test_a_0_9_14_file_is_migrated_too(self):
        """Pre-envelope files step 0 -> 1 -> 2."""
        self._write({"diarize_model": "wespeaker", "num_speakers": 0},
                    schema=0)
        s = T.current_settings()
        self.assertEqual(s["diarize_model"], "titanet")
        self.assertEqual(s["num_speakers"], 2)

    def test_choices_made_after_the_upgrade_stick(self):
        """Choosing WeSpeaker, or auto, once on the new schema is a
        decision, and must not be migrated away on the next load."""
        T._settings_save(T.validate_settings(
            {"diarize_model": "wespeaker", "num_speakers": 0,
             "diarize_threshold": 0.5}))
        s = T.current_settings()
        self.assertEqual(s["diarize_model"], "wespeaker")
        self.assertEqual(s["num_speakers"], 0)
        self.assertEqual(s["diarize_threshold"], 0.5)


class TestTranscriptionDefaults(unittest.TestCase):
    """Conditioning on the previous window off, paragraphs at a 1.0 s pause. Carried to saved
    settings by the schema 2 -> 3 migration, for untouched values only."""

    def setUp(self):
        self.addCleanup(lambda: T._settings_file().unlink(missing_ok=True))

    def _write(self, data, schema=2):
        import json
        doc = data if schema == 0 else {
            "schema": schema, "kind": "settings", "data": data}
        T._settings_file().write_text(json.dumps(doc), encoding="utf-8")

    def test_fresh_install_defaults(self):
        d = T.validate_settings({})
        self.assertIs(d["condition_on_previous_text"], False)
        self.assertEqual(d["gap"], 1.0)

    def test_old_defaults_move_to_the_new_ones(self):
        self._write({"condition_on_previous_text": True, "gap": 1.5,
                     "model": "small.en"})
        s = T.current_settings()
        self.assertIs(s["condition_on_previous_text"], False)
        self.assertEqual(s["gap"], 1.0)
        self.assertEqual(s["model"], "small.en")    # nothing else moves

    def test_a_gap_somebody_chose_is_left_alone(self):
        self._write({"gap": 2.5})
        self.assertEqual(T.current_settings()["gap"], 2.5)

    def test_a_0_9_14_file_steps_all_the_way(self):
        """Pre-envelope files step 0 -> 1 -> 2 -> 3."""
        self._write({"condition_on_previous_text": True, "gap": 1.5,
                     "num_speakers": 0}, schema=0)
        s = T.current_settings()
        self.assertIs(s["condition_on_previous_text"], False)
        self.assertEqual(s["gap"], 1.0)
        self.assertEqual(s["num_speakers"], 2)

    def test_turning_it_back_on_after_the_upgrade_sticks(self):
        T._settings_save(T.validate_settings(
            {"condition_on_previous_text": True, "gap": 1.5}))
        s = T.current_settings()
        self.assertIs(s["condition_on_previous_text"], True)
        self.assertEqual(s["gap"], 1.5)


class TestRetiredVoiceModels(unittest.TestCase):
    """Withdrawn models' weights are reclaimed: nothing else would, as
    the Models page does not list diarization models."""

    def setUp(self):
        from transcribr import diarize as D
        self.D = D
        self.dir = D._diarize_models_dir()
        self.dir.mkdir(parents=True, exist_ok=True)
        self.retired = self.dir / "campplus-voxceleb.onnx"
        self.kept = self.dir / "titanet-small.onnx"
        for f in (self.retired, self.kept):
            f.write_bytes(b"weights")
            self.addCleanup(f.unlink, missing_ok=True)

    def test_retired_weights_are_removed_and_current_ones_kept(self):
        self.D._remove_retired_models()
        self.assertFalse(self.retired.exists())
        self.assertTrue(self.kept.exists())

    def test_nothing_to_remove_is_not_an_error(self):
        self.retired.unlink()
        self.D._remove_retired_models()

    def test_every_run_reclaims_them_even_without_sherpa(self):
        """Wired into _run_diarization, which runs every time - not
        ensure_diarize_models, which runs only when a model is missing.
        sherpa_onnx is blocked so this behaves the same whether or not
        it is installed."""
        from unittest import mock
        with mock.patch.dict(sys.modules, {"sherpa_onnx": None}):
            with self.assertRaises(T.DiarizationUnavailable):
                self.D._run_diarization(None, 2, None)
        self.assertFalse(self.retired.exists())
        self.assertTrue(self.kept.exists())


# =====================================================================
# Writers and round trips
# =====================================================================

def _sample_paragraphs():
    return [
        [(0.0, 2.0, "Good morning, your Honour.")],
        [(2.0, 4.0, "Good morning, counsel.")],
        [(4.0, 6.0, "May it please the court.")],
    ]


class TestTxtRoundTrip(unittest.TestCase):
    def test_speakers_and_title_survive(self):
        speakers = ["MR SMITH", "HIS HONOUR", "MR SMITH"]
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "t.transcript.txt"
            T.write_paragraphs_to_file(
                _sample_paragraphs(), out, show_timestamp=True,
                title="R v Example", output_format="txt",
                speakers=speakers)
            parsed = T.read_paragraphs_from_file(out)
        self.assertEqual(parsed["title"], "R v Example")
        self.assertTrue(parsed["show_timestamp"])
        self.assertEqual(len(parsed["paragraphs"]), 3)
        self.assertEqual(parsed["speakers"], speakers)

    def test_unattributed_marker_blocks_carryover(self):
        speakers = ["MR SMITH", None, None]
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "t.transcript.txt"
            T.write_paragraphs_to_file(
                _sample_paragraphs(), out, output_format="txt",
                speakers=speakers)
            parsed = T.read_paragraphs_from_file(out)
        self.assertEqual(parsed["speakers"], speakers)

    def test_loaded_spans_stretch_to_next_paragraph(self):
        # Files only record paragraph start times. Parsed paragraphs must
        # come back with each end stretched to the next paragraph's start,
        # so audio playback covers the full paragraph (the bug where only
        # ~1.3s snippets played). Timestamps round to whole seconds in the
        # file, so use whole-second starts.
        paras = [
            [(0.0, 2.0, "First paragraph of the morning.")],
            [(45.0, 47.0, "Second paragraph after a long stretch.")],
            [(90.0, 92.0, "Third and final paragraph.")],
        ]
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "t.transcript.txt"
            T.write_paragraphs_to_file(paras, out, output_format="txt")
            parsed = T.read_paragraphs_from_file(out)
        spans = [(p[0][0], p[-1][1]) for p in parsed["paragraphs"]]
        self.assertEqual(spans[0], (0.0, 45.0))
        self.assertEqual(spans[1], (45.0, 90.0))
        # Last paragraph keeps the placeholder span (real end unknowable);
        # playback treats it as open-ended.
        self.assertEqual(spans[2], (90.0, 91.0))


@unittest.skipUnless(
    __import__("importlib.util", fromlist=["util"]).find_spec("docx"),
    "python-docx not installed")
class TestDocxRoundTrip(unittest.TestCase):
    def test_round_trip_with_six_speakers(self):
        paras = [[(float(i), float(i) + 1, f"Statement number {i}.")]
                 for i in range(6)]
        speakers = [f"Witness {i + 1}" for i in range(6)]
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "t.transcript.docx"
            T.write_paragraphs_to_file(
                paras, out, show_timestamp=True, title="Test",
                output_format="docx", speakers=speakers)
            self.assertGreater(out.stat().st_size, 1000)
            parsed = T.read_paragraphs_from_file(out)
        self.assertEqual(len(parsed["paragraphs"]), 6)
        self.assertEqual(parsed["speakers"], speakers)

    def test_a4_page_size(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "t.docx"
            T.write_paragraphs_to_file(
                _sample_paragraphs(), out, output_format="docx")
            sec = Document(str(out)).sections[0]
            self.assertAlmostEqual(sec.page_width.cm, 21.0, places=1)
            self.assertAlmostEqual(sec.page_height.cm, 29.7, places=1)


class TestExportSpeakerColour(unittest.TestCase):
    def test_slot_maps_to_the_light_palette(self):
        self.assertEqual(T._export_speaker_colour("2"),
                         T._PALETTES["light"]["speaker_colours"]["2"])

    def test_unlabelled_has_no_colour(self):
        for empty in (None, "", 0):
            self.assertIsNone(T._export_speaker_colour(empty))

    def test_unknown_slot_has_no_colour(self):
        self.assertIsNone(T._export_speaker_colour("47"))

    def test_export_ignores_the_active_theme(self):
        """A document is printed on white, whatever the app is set to."""
        old = T.palettes._ACTIVE_THEME
        try:
            T.palettes._ACTIVE_THEME = "dark"
            self.assertEqual(T._export_speaker_colour("1"),
                             T._PALETTES["light"]["speaker_colours"]["1"])
        finally:
            T.palettes._ACTIVE_THEME = old


@unittest.skipUnless(
    __import__("importlib.util", fromlist=["util"]).find_spec("docx"),
    "python-docx not installed")
class TestDocxStyles(unittest.TestCase):
    """Everything the writer emits is carried by a style the document
    defines, so the file looks the same wherever it is opened and a
    reader can change any of it in one edit."""

    PARAS = [[(0.0, 1.0, "First.")], [(1.0, 2.0, "Second.")],
             [(2.0, 3.0, "Third.")]]
    SLOTS = ["1", "2", None]
    NAMES = ["Counsel", "Witness", None]

    def _parts(self, **kw):
        import zipfile
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "t.docx"
            T.write_paragraphs_to_file(
                self.PARAS, out, output_format="docx", title="A title",
                speakers=self.NAMES, speaker_slots=self.SLOTS, **kw)
            with zipfile.ZipFile(out) as z:
                return (z.read("word/document.xml").decode(),
                        z.read("word/styles.xml").decode())

    @staticmethod
    def _ids(styles_xml):
        return re.findall(r'w:styleId="([^"]+)"', styles_xml)

    @staticmethod
    def _applied(doc_xml):
        return re.findall(r'<w:pStyle w:val="([^"]+)"/>', doc_xml)

    # ----- the style set ------------------------------------------------

    def test_normal_pins_the_font_for_every_script(self):
        """style.font.name alone sets only the Latin range, which left
        Word picking a theme font for anything outside it."""
        _doc, styles = self._parts()
        normal = re.search(
            r'<w:style [^>]*w:styleId="Normal".*?</w:style>', styles, re.S)
        self.assertIsNotNone(normal)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            self.assertIn(f'{attr}="Courier New"', normal.group(0))

    def test_every_transcribr_style_is_defined(self):
        _doc, styles = self._parts()
        ids = self._ids(styles)
        for want in ("TranscribrTitle", "TranscribrText", "TranscribrSpeaker",
                     "TranscribrDisclaimer", "TranscribrFooter"):
            self.assertIn(want, ids)

    def test_title_and_disclaimer_use_styles_not_direct_formatting(self):
        doc, _styles = self._parts()
        applied = self._applied(doc)
        self.assertEqual(applied[0], "TranscribrTitle")
        self.assertIn("TranscribrDisclaimer", applied)

    def test_only_the_speakers_present_get_styles(self):
        """Nine slots exist; a two-speaker transcript should not put
        nine pairs of styles in front of the reader."""
        _doc, styles = self._parts()
        ids = self._ids(styles)
        self.assertIn("TranscribrText1", ids)
        self.assertIn("TranscribrText2", ids)
        for absent in ("TranscribrText3", "TranscribrSpeaker9"):
            self.assertNotIn(absent, ids)

    def test_unused_template_styles_are_pruned(self):
        """python-docx's blank document carries ~170 styles. They bury
        the ones a reader actually needs to find."""
        _doc, styles = self._parts()
        ids = self._ids(styles)
        self.assertLess(len(ids), 30)
        for junk in ("LightShading-Accent3", "BookTitle", "Heading9"):
            self.assertNotIn(junk, ids)

    def test_styles_the_document_still_needs_survive_pruning(self):
        _doc, styles = self._parts()
        ids = self._ids(styles)
        for kept in ("Normal", "Footer", "DefaultParagraphFont"):
            self.assertIn(kept, ids)

    # ----- the colour ---------------------------------------------------

    def test_colour_lives_in_the_styles_not_the_paragraphs(self):
        """The point of the exercise: a reader changes or removes a
        speaker's colour by editing one style, not by selecting
        paragraphs and clearing formatting."""
        doc, styles = self._parts(colour_speakers=True)
        self.assertNotIn("w:shd", doc)
        self.assertIn("w:shd", styles)

    def test_each_speaker_style_carries_its_review_pane_colour(self):
        _doc, styles = self._parts(colour_speakers=True)
        for slot in ("1", "2"):
            m = re.search(
                r'<w:style [^>]*w:styleId="TranscribrText%s".*?</w:style>'
                % slot, styles, re.S)
            self.assertIsNotNone(m, f"no style for slot {slot}")
            want = T._export_speaker_colour(slot).lstrip("#").upper()
            self.assertIn(f'w:fill="{want}"', m.group(0))

    def test_label_and_body_are_shaded_alike(self):
        _doc, styles = self._parts(colour_speakers=True)
        want = T._export_speaker_colour("1").lstrip("#").upper()
        for style_id in ("TranscribrText1", "TranscribrSpeaker1"):
            m = re.search(
                r'<w:style [^>]*w:styleId="%s".*?</w:style>' % style_id,
                styles, re.S)
            self.assertIn(f'w:fill="{want}"', m.group(0))

    def test_paragraphs_reference_their_speaker_style(self):
        doc, _styles = self._parts(colour_speakers=True)
        applied = self._applied(doc)
        self.assertIn("TranscribrSpeaker1", applied)
        self.assertIn("TranscribrText1", applied)
        self.assertIn("TranscribrText2", applied)

    def test_unlabelled_paragraph_uses_the_plain_style(self):
        doc, _styles = self._parts(colour_speakers=True)
        self.assertIn("TranscribrText", self._applied(doc))

    def test_speaker_styles_exist_uncoloured_too(self):
        """So a plain transcript can be colour-coded after the fact by
        giving two styles a fill, without touching a paragraph."""
        _doc, styles = self._parts()
        ids = self._ids(styles)
        self.assertIn("TranscribrText1", ids)
        m = re.search(
            r'<w:style [^>]*w:styleId="TranscribrText1".*?</w:style>',
            styles, re.S)
        self.assertNotIn("w:shd", m.group(0))

    def test_colour_survives_a_round_trip(self):
        """Shading is presentation only - re-reading must be unaffected."""
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "t.transcript.docx"
            T.write_paragraphs_to_file(
                self.PARAS, out, output_format="docx", title="A title",
                speakers=self.NAMES, speaker_slots=self.SLOTS,
                colour_speakers=True)
            parsed = T.read_paragraphs_from_file(out)
        self.assertEqual(parsed["title"], "A title")
        self.assertEqual(len(parsed["paragraphs"]), 3)
        self.assertEqual(parsed["speakers"][:2], ["Counsel", "Witness"])

    def test_pre_0_9_15_style_names_still_read(self):
        """0.9.14 wrote "SpeakerLabel"; those files must keep opening."""
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "old.transcript.docx"
            doc = Document()
            doc.styles.add_style("SpeakerLabel", WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph("Counsel", style="SpeakerLabel")
            doc.add_paragraph("[00:00]\tA question.")
            doc.save(str(out))
            parsed = T.read_paragraphs_from_file(out)
        self.assertEqual(parsed["speakers"], ["Counsel"])


@unittest.skipUnless(
    __import__("importlib.util", fromlist=["util"]).find_spec("reportlab"),
    "reportlab not installed")
class TestPdfSpeakerColour(unittest.TestCase):
    def test_coloured_pdf_is_valid_and_differs_from_plain(self):
        paras = [[(0.0, 1.0, "First.")], [(1.0, 2.0, "Second.")]]
        slots, names = ["1", "2"], ["Counsel", "Witness"]
        with tempfile.TemporaryDirectory() as d:
            plain = Path(d) / "plain.pdf"
            colour = Path(d) / "colour.pdf"
            for out, on in ((plain, False), (colour, True)):
                T.write_paragraphs_to_file(
                    paras, out, output_format="pdf", speakers=names,
                    speaker_slots=slots, colour_speakers=on)
            self.assertTrue(colour.read_bytes().startswith(b"%PDF"))
            self.assertGreater(colour.stat().st_size,
                               plain.stat().st_size)


@unittest.skipUnless(
    __import__("importlib.util", fromlist=["util"]).find_spec("reportlab"),
    "reportlab not installed")
class TestPdfWriter(unittest.TestCase):
    def test_writes_valid_pdf(self):
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "t.transcript.pdf"
            T.write_paragraphs_to_file(
                _sample_paragraphs(), out, show_timestamp=True,
                title="R v Example", output_format="pdf",
                speakers=["MR SMITH", "HIS HONOUR", None])
            blob = out.read_bytes()
        self.assertTrue(blob.startswith(b"%PDF"), "not a PDF file")
        self.assertGreater(len(blob), 1000)

    def test_xml_specials_in_body_dont_crash(self):
        paras = [[(0.0, 1.0, "Smith & Jones <Pty> Ltd \"quoted\"")]]
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "t.pdf"
            T.write_paragraphs_to_file(paras, out, output_format="pdf")
            self.assertTrue(out.read_bytes().startswith(b"%PDF"))


def _have_numpy():
    try:
        import numpy  # noqa: F401
        return True
    except ImportError:
        return False


class TestDeadAirDetection(unittest.TestCase):
    RATE = 16000

    def _samples(self, *parts):
        """Build a mono array from (seconds, amplitude) parts."""
        import numpy as np
        rng = np.random.default_rng(7)
        chunks = []
        for secs, amp in parts:
            n = int(secs * self.RATE)
            if amp == 0.0:
                chunks.append(np.zeros(n, dtype=np.float32))
            else:
                chunks.append(
                    (rng.standard_normal(n) * amp).astype(np.float32))
        return np.concatenate(chunks)

    @unittest.skipUnless(_have_numpy(), "needs numpy")
    def test_detects_leading_and_middle_silence(self):
        audio = self._samples((6, 0.0), (4, 0.1), (7, 0.0), (3, 0.1))
        spans = T.detect_silences(audio, self.RATE)
        self.assertEqual(spans, [(0.0, 6.0), (10.0, 17.0)])

    @unittest.skipUnless(_have_numpy(), "needs numpy")
    def test_trailing_silence_runs_to_end_of_file(self):
        audio = self._samples((4, 0.1), (6.3, 0.0))
        self.assertEqual(T.detect_silences(audio, self.RATE),
                         [(4.0, 10.3)])

    @unittest.skipUnless(_have_numpy(), "needs numpy")
    def test_short_gaps_and_quiet_speech_are_not_dead_air(self):
        # A 3s gap is an ordinary pause; a quiet-but-live channel
        # (soft speech, room tone) sits well above the -60 dBFS gate.
        audio = self._samples((5, 0.1), (3, 0.0), (5, 0.1))
        self.assertEqual(T.detect_silences(audio, self.RATE), [])
        audio = self._samples((6, 0.01), (6, 0.1))
        self.assertEqual(T.detect_silences(audio, self.RATE), [])

    def test_phantom_segments_inside_silence_are_dropped(self):
        spans = [(0.0, 40.0)]
        segments = [
            (2.0, 6.0, "Thanks for watching!"),
            (38.5, 43.0, "Real speech."),
            (45.0, 50.0, "More speech."),
        ]
        kept, dropped = T._drop_segments_in_silence(segments, spans)
        self.assertEqual(dropped, 1)
        self.assertEqual([s[2] for s in kept],
                         ["Real speech.", "More speech."])

    def test_markers_are_spliced_in_time_order(self):
        paragraphs = [
            [(41.0, 44.0, "Hello there.")],
            [(50.0, 53.0, "How are you?")],
        ]
        spans = [(0.0, 40.0), (45.0, 49.5)]
        merged, letters = T._splice_no_audio_markers(
            paragraphs, ["1", "2"], spans)
        self.assertEqual(
            [p[0][2] for p in merged],
            ["[No audio from 00:00 to 00:40]",
             "Hello there.",
             "[No audio from 00:45 to 00:49]",
             "How are you?"])
        self.assertEqual(letters, [None, "1", None, "2"])
        # Without diarisation the labels stay absent entirely.
        merged, letters = T._splice_no_audio_markers(
            paragraphs, None, spans)
        self.assertEqual(len(merged), 4)
        self.assertIsNone(letters)

    @unittest.skipUnless(_have_numpy(), "needs numpy")
    def test_pipeline_marks_dead_air_and_drops_phantoms(self):
        import queue as queue_mod
        audio = self._samples((10, 0.0), (10, 0.1))
        segments = [
            (1.0, 4.0, "Phantom text."),
            (11.0, 14.0, "Real speech."),
        ]
        params = {"_audio": audio, "input": "unused", "gap": 2.0}
        q = queue_mod.Queue()
        paragraphs, letters = T._paragraphs_with_speakers(
            segments, None, params, q, None)
        self.assertEqual(
            [p[0][2] for p in paragraphs],
            ["[No audio from 00:00 to 00:10]", "Real speech."])
        self.assertIsNone(letters)


class TestHallucinationDetection(unittest.TestCase):
    @staticmethod
    def _paras(*texts):
        return [[(float(i), float(i) + 1.0, t)]
                for i, t in enumerate(texts)]

    def test_flags_repeated_phrase_loop(self):
        paras = self._paras(
            "The vehicle was blue.",
            "Thanks for watching. Thanks for watching. Thanks for "
            "watching.",
            "The officer then spoke.")
        self.assertEqual(T.detect_hallucinations(paras), [1])

    def test_flags_single_word_loop(self):
        paras = self._paras("So the witness said this.",
                            "you you you you you you")
        self.assertEqual(T.detect_hallucinations(paras), [1])

    def test_flags_two_repeat_paragraph_that_is_all_loop(self):
        # A paragraph that is nothing but a phrase repeated twice, long
        # enough not to be an ordinary short repeat.
        paras = self._paras("Setting the scene here.",
                            "Thanks for watching. Thanks for watching.")
        self.assertEqual(T.detect_hallucinations(paras), [1])

    def test_short_two_repeat_phrase_not_flagged(self):
        # "Thank you. Thank you." is only four words - ordinary, not a
        # runaway loop.
        self.assertEqual(
            T.detect_hallucinations(self._paras("Thank you. Thank you.")),
            [])

    def test_flags_consecutive_duplicate_paragraphs(self):
        paras = self._paras(
            "Opening remarks by counsel.",
            "The signal was lost at this point.",
            "The signal was lost at this point.",
            "Cross-examination resumed.")
        self.assertEqual(T.detect_hallucinations(paras), [1, 2])

    def test_ignores_ordinary_courtroom_repetition(self):
        # Short affirmations and a phrase used twice non-consecutively
        # are not loops.
        paras = self._paras(
            "Yes. Yes.",
            "No, no.",
            "I went to the shop and the shop was closed.",
            "That is correct.")
        self.assertEqual(T.detect_hallucinations(paras), [])

    def test_two_word_short_paragraph_not_flagged(self):
        # Guards the coverage floor: a 2-word paragraph shouldn't trip.
        self.assertEqual(
            T.detect_hallucinations(self._paras("Thank you.")), [])

    def test_low_confidence_paragraph_flagged(self):
        # The engine's own confidence is a second hint: a paragraph
        # most of whose words scored very low is flagged even without
        # any repetition.
        paras = self._paras("A clearly spoken sentence here.",
                            "garbled indistinct crosstalk mumble noise",
                            "Another clear sentence follows on.")
        confidences = [[0.9, 0.9, 0.9, 0.9, 0.9],
                       [0.1, 0.2, 0.15, 0.3, 0.25],
                       [0.9, 0.9, 0.9, 0.9]]
        self.assertEqual(
            T.detect_hallucinations(paras, confidences=confidences), [1])
        # Without the confidence hint the text alone reads fine.
        self.assertEqual(T.detect_hallucinations(paras), [])

    def test_soft_but_confident_passage_not_flagged(self):
        # Medium confidence (a soft-spoken but genuine passage) is left
        # alone - only the strong low band counts.
        paras = self._paras("A quietly spoken but real sentence.")
        self.assertEqual(
            T.detect_hallucinations(paras, confidences=[[0.5, 0.5, 0.5,
                                                         0.5, 0.5]]), [])


class TestUnknownFormatFallsBackToTxt(unittest.TestCase):
    def test_unknown_format_writes_text(self):
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "t.weird"
            T.write_paragraphs_to_file(
                _sample_paragraphs(), out, output_format="weird")
            self.assertIn("Good morning", out.read_text(encoding="utf-8"))



if __name__ == "__main__":
    unittest.main(verbosity=2)


# =====================================================================
# The harness itself
# =====================================================================

class TestTheSuiteLeavesTheRealConfigAlone(unittest.TestCase):
    """A regression test for the suite, not the app.

    Review autosave is a debounced timer. Tests that build a session
    and never close it leave one armed, and it used to fire after the
    module teardown had put the config directory back - writing fixture
    transcripts into the user's real autosave.json, which the app would
    then offer to restore."""

    def test_the_config_dir_is_redirected(self):
        self.assertEqual(str(T._config_dir()), _harness.CONFIG_DIR)

    def test_the_redirect_is_not_undone_between_modules(self):
        self.assertEqual(os.environ.get("TRANSCRIBR_CONFIG_DIR"),
                         _harness.CONFIG_DIR)

    def test_the_log_is_redirected_too(self):
        """Tests exercise the error paths deliberately; their
        tracebacks must not land in the log a user would send in."""
        self.assertEqual(str(T._log_dir()), _harness.LOG_DIR)
        T._log("probe from the suite")
        self.assertTrue((Path(_harness.LOG_DIR) / "transcribr.log").exists())

    def test_a_late_autosave_timer_cannot_reach_the_real_directory(self):
        """The exact shape of the leak: a session nobody closed, whose
        timer fires long after the test that made it has finished."""
        real = _harness.real_config_dir() / "autosave.json"
        before = real.stat().st_mtime if real.exists() else None

        model = T.TranscriptModel(
            [[(0.0, 1.0, "Fixture text.")]],
            on_autosave=lambda paras, spk, names: T._autosave_save({
                "out_path": "/tmp/leak-probe.txt",
                "paragraphs": [[list(s) for s in p] for p in paras],
                "speakers": list(spk), "speaker_names": dict(names),
            }))
        model.set_speaker(0, "1")
        model.flush_autosave()          # what the stray timer would do

        written = Path(_harness.CONFIG_DIR) / "autosave.json"
        self.assertTrue(written.exists(), "autosave did not go to the temp dir")
        self.assertEqual(T._autosave_load()["out_path"], "/tmp/leak-probe.txt")
        T._autosave_clear()

        after = real.stat().st_mtime if real.exists() else None
        self.assertEqual(before, after,
                         "the suite touched the real autosave.json")
        model.close()
