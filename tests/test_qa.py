"""Q&A layout, numbered like an ERISP (transcribr/qa.py)."""

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import _harness  # noqa: F401

import transcribr as T
from transcribr import qa as QA


def para(t0, text):
    return [(t0, t0 + 2.0, text)]


INTERVIEW = [
    para(0, "What is your full name?"),
    para(3, "John Citizen."),
    para(6, "Where were you on Friday night?"),
    para(9, "I was at home."),
    para(12, "All night?"),
    para(15, "Yes, all night."),
]


class TestMarks(unittest.TestCase):

    def test_questions_are_numbered_and_answers_marked(self):
        self.assertEqual(QA.qa_marks(["Q", "A", "Q", "A"]),
                         ["Q1", "A", "Q2", "A"])

    def test_a_side_carrying_on_is_not_marked_again(self):
        """A question over two paragraphs is one question; an answer
        over two paragraphs is one answer."""
        self.assertEqual(QA.qa_marks(["Q", "Q", "A", "A", "Q"]),
                         ["Q1", "", "A", "", "Q2"])

    def test_unplaced_paragraphs_carry_no_mark(self):
        self.assertEqual(QA.qa_marks([None, "Q", None, "Q"]),
                         ["", "Q1", "", "Q2"])


class TestRoles(unittest.TestCase):

    def test_labels_decide_when_every_paragraph_has_one(self):
        slots = ["1", "2", "1", "2", "1", "2"]
        roles, askers, method = QA.qa_roles(INTERVIEW, slots)
        self.assertEqual(askers, ["1"])
        self.assertEqual(roles, ["Q", "A", "Q", "A", "Q", "A"])
        self.assertEqual(method, "labels")

    def test_named_askers_override_the_suggestion(self):
        slots = ["1", "2", "1", "2", "1", "2"]
        roles, askers, _ = QA.qa_roles(INTERVIEW, slots, askers=["2"])
        self.assertEqual(roles, ["A", "Q", "A", "Q", "A", "Q"])

    def test_no_labels_falls_back_to_the_wording(self):
        roles, askers, method = QA.qa_roles(INTERVIEW, [None] * 6)
        self.assertEqual(method, "wording")
        self.assertEqual(roles, ["Q", "A", "Q", "A", "Q", "A"])

    def test_wording_works_without_pause_timings(self):
        """A re-opened transcript knows only start times: its paragraphs
        butt up against each other, so there are no pauses to go on."""
        tight = [[(3.0 * i, 3.0 * (i + 1), p[0][2])]
                 for i, p in enumerate(INTERVIEW)]
        roles, _, _ = QA.qa_roles(tight, [None] * 6)
        self.assertEqual(roles, ["Q", "A", "Q", "A", "Q", "A"])

    def test_one_labelled_speaker_cannot_tell_sides_apart(self):
        self.assertEqual(QA.suggest_askers(INTERVIEW, ["1"] * 6), [])

    def test_the_note_warns_when_the_wording_decided(self):
        self.assertIn("Check each against the recording",
                      QA.method_note("wording", []))
        self.assertIn("questions by Det Smith",
                      QA.method_note("labels", ["Det Smith"]))


class TestWriting(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="transcribr-qa-")
        self.addCleanup(self.tmp.cleanup)

    def test_text_layout(self):
        out = Path(self.tmp.name) / "x.qa.txt"
        QA.write_qa(INTERVIEW[:2], out, ["Q1", "A"], fmt="txt",
                    title="ERISP", note="A note.")
        lines = out.read_text(encoding="utf-8").splitlines()
        self.assertEqual(lines[:6], ["ERISP", "", "A note.", "",
                                     "Q1    What is your full name?",
                                     "A     John Citizen."])

    def test_word_layout(self):
        try:
            import docx
        except ImportError:
            self.skipTest("needs python-docx")
        out = Path(self.tmp.name) / "x.qa.docx"
        QA.write_qa(INTERVIEW[:3], out, ["Q1", "A", "Q2"], fmt="docx",
                    note="A note.")
        texts = [p.text for p in docx.Document(str(out)).paragraphs]
        self.assertIn("Q1\tWhat is your full name?", texts)
        self.assertIn("A\tJohn Citizen.", texts)
        self.assertIn("A note.", texts)


class TestExport(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="transcribr-qa-")
        self.addCleanup(self.tmp.cleanup)

    def test_review_session_exports_qa_beside_the_transcript(self):
        try:
            import docx
        except ImportError:
            self.skipTest("needs python-docx")
        out = str(Path(self.tmp.name) / "int.transcript.txt")
        info = {"paragraphs": INTERVIEW, "out_path": out,
                "show_timestamp": False, "title": "Interview",
                "output_format": "txt", "result": {"segments": []},
                "extra_formats": [], "loaded": False, "audio_path": None,
                "word_conf": None}
        session = T.ReviewSession.from_fresh(info, T.EventBroker())
        path = session.export(session.model.rev, "pdf", layout="qa")
        self.assertTrue(path.endswith("int.transcript.qa.docx"))
        texts = [p.text for p in docx.Document(path).paragraphs]
        self.assertIn("Q3\tAll night?", texts)
        self.assertTrue(any("from the wording" in t for t in texts))
        self.assertFalse(session.closed)


TWO_OFFICERS = [
    para(0, "What is your full name?"),          # 1 asks
    para(3, "John Citizen."),                    # 3 answers
    para(6, "[No audio from 0:08 to 0:20]"),     # marker
    para(20, "Where were you on Friday night?"),  # 1 asks
    para(23, "At home."),                        # 3
    para(26, "Who else was there?"),             # 2 takes over
    para(29, "Nobody."),                         # 3
    para(32, "And on Saturday?"),                # 2
    para(35, "Can I add something?"),            # 1 straight after 2
    para(38, "At my sister's."),                 # 3
]
TWO_SLOTS = ["1", "3", None, "1", "3", "2", "3", "2", "1", "3"]
NAMES = {"1": "Det Smith", "2": "Det Jones", "3": "John Citizen"}


class TestNoAudioAndAskers(unittest.TestCase):

    def test_no_audio_markers_get_no_role_or_number(self):
        lay = QA.qa_layout(TWO_OFFICERS, TWO_SLOTS, askers=["1", "2"],
                           names=NAMES)
        row = lay["rows"][2]
        self.assertEqual((row["role"], row["mark"]), (None, ""))
        # ...and the marker does not count against "labels".
        self.assertEqual(lay["method"], "labels")

    def test_no_audio_markers_are_skipped_by_the_wording_model(self):
        roles, _, _ = QA.qa_roles(TWO_OFFICERS[:5], [None] * 5)
        self.assertIsNone(roles[2])

    def test_a_change_of_asker_is_named_on_its_first_question(self):
        lay = QA.qa_layout(TWO_OFFICERS, TWO_SLOTS, askers=["1", "2"],
                           names=NAMES)
        marks = [r["mark"] for r in lay["rows"]]
        askers = [r["asker"] for r in lay["rows"]]
        self.assertEqual(marks, ["Q1", "A", "", "Q2", "A", "Q3", "A",
                                 "Q4", "Q5", "A"])
        self.assertEqual(askers, ["Det Smith", None, None, None, None,
                                  "Det Jones", None, None, "Det Smith",
                                  None])

    def test_one_asker_is_never_named(self):
        lay = QA.qa_layout(INTERVIEW, ["1", "2"] * 3, names=NAMES)
        self.assertEqual([r["asker"] for r in lay["rows"]], [None] * 6)


class TestQaMode(unittest.TestCase):
    """Q&A as a mode of the review session: shown, saved, reopened."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="transcribr-qa-")
        self.addCleanup(self.tmp.cleanup)

    def _session(self, fmt="docx"):
        out = str(Path(self.tmp.name) / f"int.transcript.{fmt}")
        info = {"paragraphs": TWO_OFFICERS, "out_path": out,
                "show_timestamp": True, "title": "Interview",
                "output_format": fmt, "result": {"segments": []},
                "extra_formats": [], "loaded": False, "audio_path": None,
                "word_conf": None, "preset_speakers": TWO_SLOTS,
                "preset_speaker_names": NAMES}
        return T.ReviewSession.from_fresh(info, T.EventBroker())

    def test_switching_mode_shows_the_marks(self):
        s = self._session()
        self.assertIsNone(s.payload()["qa"])
        full = s.mutate(s.model.rev, "layout", {"layout": "qa"})
        self.assertEqual(full["layout"], "qa")
        self.assertTrue(full["qa"]["auto"])
        full = s.mutate(full["rev"], "askers", {"askers": ["1", "2"]})
        self.assertEqual(full["qa"]["askers"], ["1", "2"])
        self.assertFalse(full["qa"]["auto"])
        self.assertEqual(full["qa"]["rows"][5]["asker"], "Det Jones")
        # A speaker change replies slim, and still carries the marks.
        delta = s.mutate(full["rev"], "speaker", {"index": 5, "slot": "1"})
        self.assertIsNone(delta["qa"]["rows"][5]["asker"])
        full = s.mutate(delta["rev"], "layout", {"layout": "plain"})
        self.assertIsNone(full["qa"])

    def test_docx_saved_in_qa_mode_reopens_in_qa_mode(self):
        try:
            import docx
        except ImportError:
            self.skipTest("needs python-docx")
        s = self._session("docx")
        s.mutate(s.model.rev, "layout", {"layout": "qa"})
        s.mutate(s.model.rev, "askers", {"askers": ["1", "2"]})
        out = s.save(s.model.rev, "labels")
        texts = [p.text for p in docx.Document(out).paragraphs]
        self.assertTrue(any(t.startswith("Q3\t") and "(Det Jones) Who"
                            in t for t in texts))
        self.assertTrue(any(t.startswith("\t") and "[No audio" in t
                            for t in texts))

        info = T.open_transcript_info(out)
        self.assertEqual(info["layout"], "qa")
        names = info["preset_speaker_names"]
        speakers = [names.get(x) if x else None
                    for x in info["preset_speakers"]]
        self.assertEqual(speakers, [NAMES.get(x) if x else None
                                    for x in TWO_SLOTS])
        self.assertEqual(sorted(names[a] for a in info["askers"]),
                         ["Det Jones", "Det Smith"])
        bodies = [p[0][2] for p in info["paragraphs"]]
        self.assertEqual(bodies, [p[0][2] for p in TWO_OFFICERS])

    def test_txt_saved_in_qa_mode_reopens_by_side(self):
        s = self._session("txt")
        s.mutate(s.model.rev, "layout", {"layout": "qa"})
        s.mutate(s.model.rev, "askers", {"askers": ["1", "2"]})
        out = s.save(s.model.rev, "labels")
        info = T.open_transcript_info(out)
        self.assertEqual(info["layout"], "qa")
        names = info["preset_speaker_names"]
        speakers = [names.get(x) if x else None
                    for x in info["preset_speakers"]]
        self.assertEqual(speakers, [
            "Det Smith", "Answerer", None, "Det Smith", "Answerer",
            "Det Jones", "Answerer", "Det Jones", "Det Smith",
            "Answerer"])
        bodies = [p[0][2] for p in info["paragraphs"]]
        self.assertEqual(bodies, [p[0][2] for p in TWO_OFFICERS])

    def test_pdf_saves_in_qa_mode(self):
        try:
            import reportlab  # noqa: F401
        except ImportError:
            self.skipTest("needs reportlab")
        s = self._session("pdf")
        s.mutate(s.model.rev, "layout", {"layout": "qa"})
        out = s.save(s.model.rev, "labels")
        self.assertTrue(Path(out).read_bytes().startswith(b"%PDF"))

if __name__ == "__main__":
    unittest.main()
