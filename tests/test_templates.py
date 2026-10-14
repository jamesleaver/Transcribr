"""Word templates: plain, court and ERISP (writers._docx_apply_template)."""

import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import _harness  # noqa: F401

import transcribr as T
from transcribr import qa as QA

try:
    import docx
    from docx.oxml.ns import qn
except ImportError:          # pragma: no cover
    docx = None

PARAS = [[(0.0, 4.0, "What is your full name?")],
         [(5.0, 9.0, "John Citizen.")],
         [(10.0, 14.0, "Where were you on Friday?")],
         [(15.0, 19.0, "At home, all night.")]]
SPEAKERS = ["Det Smith", "John Citizen", "Det Smith", "John Citizen"]


@unittest.skipIf(docx is None, "needs python-docx")
class TestTemplates(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="transcribr-tpl-")
        self.addCleanup(self.tmp.cleanup)

    def _write(self, template):
        out = Path(self.tmp.name) / f"{template}.docx"
        T.write_paragraphs_to_file(
            PARAS, out, show_timestamp=True, title="R v Citizen",
            output_format="docx", speakers=SPEAKERS,
            audio_path="/recordings/int-0042.m4a", template=template)
        return out

    def _ln(self, d):
        return d.sections[0]._sectPr.find(qn("w:lnNumType"))

    def test_plain_is_unchanged(self):
        d = docx.Document(str(self._write("plain")))
        self.assertIsNone(self._ln(d))
        self.assertEqual(d.styles["Normal"].font.name, T.DOCX_FONT)

    def test_court_numbers_lines_per_page_and_heads_each_page(self):
        d = docx.Document(str(self._write("court")))
        ln = self._ln(d)
        self.assertIsNotNone(ln)
        self.assertEqual(ln.get(qn("w:restart")), "newPage")
        self.assertEqual(ln.get(qn("w:countBy")), "1")
        self.assertEqual(d.styles["Normal"].font.name, "Times New Roman")
        self.assertEqual(d.styles["Normal"].paragraph_format.line_spacing,
                         2.0)
        head = d.sections[0].header.paragraphs[0].text
        self.assertEqual(head, "R v Citizen\tint-0042.m4a")

    def test_erisp_carries_the_record_of_interview_heading(self):
        d = docx.Document(str(self._write("erisp")))
        self.assertIsNotNone(self._ln(d))
        heads = [p.text for p in d.sections[0].header.paragraphs]
        self.assertEqual(heads, ["RECORD OF INTERVIEW", "R v Citizen"])

    def test_every_template_reopens_the_same(self):
        for template in T.DOC_TEMPLATES:
            with self.subTest(template=template):
                parsed = T.read_paragraphs_from_file(
                    str(self._write(template)))
                self.assertEqual(parsed["title"], "R v Citizen")
                self.assertEqual(parsed["speakers"], SPEAKERS)
                self.assertEqual([p[0][2] for p in parsed["paragraphs"]],
                                 [p[0][2] for p in PARAS])

    def test_qa_layout_takes_a_template_and_still_reopens(self):
        out = Path(self.tmp.name) / "qa.docx"
        lay = QA.qa_layout(PARAS, ["1", "2", "1", "2"],
                           names={"1": "Det Smith", "2": "John Citizen"})
        QA.write_qa(PARAS, out, [r["mark"] for r in lay["rows"]],
                    fmt="docx", title="R v Citizen", template="erisp",
                    meta={"v": 1, "speakers": SPEAKERS,
                          "askers": ["Det Smith"], "named": []})
        d = docx.Document(str(out))
        self.assertIsNotNone(self._ln(d))
        parsed = T.read_paragraphs_from_file(str(out))
        self.assertEqual(parsed["layout"], "qa")
        self.assertEqual(parsed["speakers"], SPEAKERS)


if __name__ == "__main__":
    unittest.main()
